"""Build the curated final-poll table from Mark Pack's PollBase workbook.

PollBase (https://www.markpack.org.uk/opinion-polls/) lists every national
voting-intention poll since 1943, one sheet per parliament, each sheet ending
in the general election's exit polls and a ``Result`` row (GB vote shares).
This script picks each election's final polls, averages them and writes them
beside the result to ``data/models/westminster/uk_final_poll_errors.json``,
the input of ``scripts/calibrate_uncertainty.py``'s national polling error.

**Final-poll rule** (applied uniformly to every election):

1. A poll row has a fieldwork start date and a numeric Con, Lab and LD share.
2. Its fieldwork end (the last day number in the ``Fieldwork`` text, read
   forward from the start date, so ``"31-1"`` in May ends on 1 June) falls in
   the last ``WINDOW_DAYS`` days before polling day or on polling day itself,
   and its fieldwork started before polling day (no on-the-day polls).
3. It was published by polling day: its ``Published`` cell (a day of the
   month read forward from the fieldwork start, or a full date) is on or
   before polling day. An empty ``Published`` cell does not count as
   published; PollBase leaves it empty for polls released only later (e.g.
   Survation's withheld 6 May 2015 phone poll), and later releases carry a
   full date (Focaldata's 2024 poll: 14 April 2025).
4. Of each pollster's qualifying polls, only the last counts: latest
   fieldwork end, then latest publication, then the later row.
5. A party's average is the equal-weight mean over the pollsters that report
   it (a blank cell is "not reported", not 0).

Parties: Con, Lab and LD always; Green, UKIP, the Brexit Party and Reform UK
where both the election's ``Result`` row and at least one final poll report
them. PollBase has no SNP column, so the SNP never appears.

Shares are used as published. A few pollsters (e.g. Savanta and Survation in
2024) report UK-wide shares, which put a GB party about 2.5% (relative)
lower than a GB share; their ``geography`` is recorded on each poll.

Usage:
    python data/scripts/build_final_poll_errors.py PATH/TO/PollBase.xlsx
    python data/scripts/build_final_poll_errors.py PATH/TO/PollBase.xlsx --dry-run
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
import zlib
from collections.abc import Iterable, Iterator, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import fmean
from xml.etree.ElementTree import ParseError

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

SCRIPT_DIR = Path(__file__).resolve().parent
DATA_DIR = SCRIPT_DIR.parent
if str(DATA_DIR) not in sys.path:
    sys.path.insert(0, str(DATA_DIR))

from model_support.calibration import (
    FINAL_POLL_ERRORS_PATH,
    FINAL_POLL_SCHEMA,
    POLLBASE_CREDIT,
    POLLBASE_URL,
)
from model_support.io import write_text_files

WINDOW_DAYS = 3

# Polling days of the general elections calibrated against (1992-2024).
POLLING_DAYS: tuple[date, ...] = (
    date(1992, 4, 9),
    date(1997, 5, 1),
    date(2001, 6, 7),
    date(2005, 5, 5),
    date(2010, 5, 6),
    date(2015, 5, 7),
    date(2017, 6, 8),
    date(2019, 12, 12),
    date(2024, 7, 4),
)

# PollBase column header -> (party key, label). The key is what the
# calibration and its cross-check against the database read.
PARTY_COLUMNS: Mapping[str, tuple[str, str]] = {
    "Con": ("con", "Conservative"),
    "Lab": ("lab", "Labour"),
    "LD": ("ld", "Liberal Democrats"),
    "Green": ("green", "Green"),
    "UKIP": ("ukip", "UKIP"),
    "Ukip": ("ukip", "UKIP"),
    "BXP": ("brexit", "Brexit Party"),
    "BXP/Reform": ("reform", "Reform UK"),
    "Reform": ("reform", "Reform UK"),
}
REQUIRED_PARTIES: tuple[str, ...] = ("con", "lab", "ld")

RULE_TEXT = (
    "For each election, each pollster's last poll (latest fieldwork end, then "
    "latest publication, then later row) whose fieldwork ended in the "
    f"{WINDOW_DAYS} days before polling day or on it, started before polling "
    "day, and was published on or before polling day (a blank Published cell "
    "does not count); a party's average is the equal-weight mean over the "
    "pollsters reporting it. Results are PollBase's Result rows (GB vote "
    "shares). Shares as published, UK-basis polls included."
)

RESULT_LABEL = "Result"
# What a damaged or non-Excel file raises from openpyxl, zipfile and zlib.
_CORRUPT_WORKBOOK = (
    zipfile.BadZipFile,
    InvalidFileException,
    zlib.error,
    EOFError,
    KeyError,
    ValueError,
    ParseError,
)
RATINGS_HEADER = "Leader:"
GEOGRAPHIES = frozenset({"UK", "GB"})
# A sheet stops after this many consecutive empty rows (two sheets are
# padded out to Excel's last row).
MAX_EMPTY_RUN = 200
# A Result row belongs to the election whose polls precede it in its sheet:
# the sheet's last poll started at most this long before polling day.
RESULT_LOOKBACK = timedelta(days=30)
MAX_FIELDWORK_DAYS = 31
# Polls are parsed only when they started this close before a polling day.
READ_LOOKBACK = timedelta(days=45)

# PollBase rows whose dates contradict each other, skipped by name. Each entry
# is (sheet, row, pollster, fieldwork text) -> why; a listed row that no longer
# matches (a new edition renumbered it) stops the build for review.
SKIPPED_ROWS: Mapping[tuple[str, int, str, str], str] = {
    ("05-10", 624, "YouGov", "16"): (
        "fieldwork text '16' (April) contradicts its start date, 17 April 2010; "
        "19 days before polling day, outside the final-poll window either way"
    ),
    ("19-24", 1894, "YouGov", "11-18"): (
        "YouGov/Sky MRP, fieldwork 11-18 June 2024, printed as published on the "
        "16th, before its fieldwork ended; outside the final-poll window either way"
    ),
}

_FIELDWORK_RE = re.compile(
    r"^(?:c\.\s*)?(?:(\d{1,2})|\?)(?:\s*[-/]\s*(\d{1,2}))?\s*\??$"
)
_DAY_MONTH_RANGE_RE = re.compile(
    r"^\d{1,2}(?:/\d{1,2})?\s*-\s*(\d{1,2})/(\d{1,2})$"
)
_PUBLISHED_RE = re.compile(r"^(\d{1,2})(?:\s*[-/]\s*\d{1,2})?$")


class PollBaseError(ValueError):
    """The workbook does not have the layout this builder reads."""


# ── Reading the workbook ──────────────────────────────────────────────────────


@dataclass(frozen=True, slots=True)
class PollRow:
    """One voting-intention row of a PollBase sheet.

    Attributes:
        sheet: The sheet's title.
        row: 1-based row number in the sheet.
        pollster: The ``Polling`` cell, stripped.
        publisher: The ``Publisher`` cell, stripped (None when blank).
        fieldwork: The ``Fieldwork`` text as printed (None when blank).
        start: Fieldwork start date.
        end: Fieldwork end date (:func:`fieldwork_end`).
        published: Publication date (:func:`published_date`), None when
            the cell is blank.
        geography: "UK" or "GB" when the sheet's geography cell says so
            (the column sometimes holds a methodology note instead).
        shares: Party key -> share (percent) for the party columns with a
            number.
    """

    sheet: str
    row: int
    pollster: str
    publisher: str | None
    fieldwork: str | None
    start: date
    end: date
    published: date | None
    geography: str | None
    shares: Mapping[str, float]


@dataclass(frozen=True, slots=True)
class ResultRow:
    """A ``Result`` row: the GB vote shares of the election it follows.

    Attributes:
        sheet: The sheet's title.
        row: 1-based row number.
        day: The ``Published`` cell (polling day's day of the month).
        last_poll_start: Latest poll fieldwork start above it in the sheet.
        shares: Party key -> share (percent).
    """

    sheet: str
    row: int
    day: int | None
    last_poll_start: date | None
    shares: Mapping[str, float]


@dataclass(slots=True)
class Workbook:
    """Every poll and result row read from the workbook.

    Attributes:
        polls: The poll rows near an election.
        results: Every ``Result`` row.
        skipped: The skip-list rows found (and skipped), with the reasons.
    """

    polls: list[PollRow] = field(default_factory=list)
    results: list[ResultRow] = field(default_factory=list)
    skipped: dict[tuple[str, int, str, str], str] = field(default_factory=dict)


def _day_on_or_after(start: date, day: int) -> date:
    """Return the first date on or after *start* whose day of month is *day*."""
    for offset in range(MAX_FIELDWORK_DAYS + 31):
        candidate = start + timedelta(days=offset)
        if candidate.day == day:
            return candidate
    raise PollBaseError(f"no day {day} within two months of {start}")


def fieldwork_end(text: object, start: date) -> date:
    """Return the fieldwork end date from PollBase's ``Fieldwork`` cell.

    The cell holds day numbers only: ``"24"``, ``"24-25"``, ``"31-1"`` (into
    the next month), sometimes marked approximate (``"c.6-9"``) or uncertain
    (``"25-28?"``, ``"?-9"``), or a bare number; the end is the last day
    number, read forward from *start*. A long poll may print its end as a
    day/month pair (``"30/5-21/6"``, ``"10-1/7"``), which is that date.

    Args:
        text: The cell value.
        start: The fieldwork start date (the next column).

    Returns:
        The end date; *start* when the cell is blank.

    Raises:
        PollBaseError: The text is not a day or day range, or the range runs
            longer than ``MAX_FIELDWORK_DAYS``.
    """
    if text is None or (isinstance(text, str) and not text.strip()):
        return start
    if isinstance(text, bool) or not isinstance(text, int | float | str):
        raise PollBaseError(f"unreadable fieldwork {text!r}")
    if isinstance(text, int | float):
        if not float(text).is_integer():
            raise PollBaseError(f"unreadable fieldwork {text!r}")
        end_day = int(text)
    elif (dated := _DAY_MONTH_RANGE_RE.match(text.strip())) is not None:
        return _dated_end(text, start, int(dated.group(1)), int(dated.group(2)))
    else:
        match = _FIELDWORK_RE.match(text.strip())
        last = None if match is None else (match.group(2) or match.group(1))
        if last is None:
            raise PollBaseError(f"unreadable fieldwork {text!r}")
        end_day = int(last)
    if not 1 <= end_day <= 31:
        raise PollBaseError(f"unreadable fieldwork {text!r}")
    end = _day_on_or_after(start, end_day)
    if (end - start).days > MAX_FIELDWORK_DAYS:
        raise PollBaseError(f"fieldwork {text!r} from {start} is too long")
    return end


def _dated_end(text: str, start: date, day: int, month: int) -> date:
    try:
        end = date(start.year + (month < start.month), month, day)
    except ValueError as err:
        raise PollBaseError(f"unreadable fieldwork {text!r}") from err
    if not 0 <= (end - start).days <= MAX_FIELDWORK_DAYS:
        raise PollBaseError(f"fieldwork {text!r} from {start} is out of range")
    return end


def published_date(value: object, start: date) -> date | None:
    """Return the publication date from PollBase's ``Published`` cell.

    Args:
        value: The cell: a day of the month (``5``, ``"26/27"`` and
            ``"2-3"`` for rolling polls, read as the first day), or a full
            date for polls released long after fieldwork.
        start: The fieldwork start date; a day number is read forward from
            it.

    Returns:
        The date, or None when the cell is blank.

    Raises:
        PollBaseError: The cell is neither a day nor a date.
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, bool):
        raise PollBaseError(f"unreadable published date {value!r}")
    if isinstance(value, int | float) and float(value).is_integer():
        day = int(value)
    elif isinstance(value, str) and (match := _PUBLISHED_RE.match(value.strip())):
        day = int(match.group(1))
    else:
        raise PollBaseError(f"unreadable published date {value!r}")
    if not 1 <= day <= 31:
        raise PollBaseError(f"unreadable published date {value!r}")
    return _day_on_or_after(start, day)


def _text(value: object) -> str | None:
    if value is None:
        return None
    stripped = str(value).strip()
    return stripped or None


def _geography(value: object) -> str | None:
    text = _text(value)
    return text if text in GEOGRAPHIES else None


def _number(value: object) -> float | None:
    if isinstance(value, bool) or not isinstance(value, int | float):
        return None
    return float(value)


@dataclass(frozen=True, slots=True)
class _Columns:
    fieldwork: int
    start: int
    published: int
    polling: int
    publisher: int
    geography: int | None
    parties: Mapping[str, int]


def _columns(sheet: str, header: Sequence[object]) -> _Columns | None:
    """Locate a sheet's columns by header, or None for a non-poll sheet."""
    names = [_text(cell) for cell in header]
    if "Fieldwork" not in names or "Polling" not in names:
        return None
    try:
        fieldwork = names.index("Fieldwork")
        published = names.index("Published")
        polling = names.index("Polling")
        publisher = names.index("Publisher")
    except ValueError as err:
        raise PollBaseError(f"sheet {sheet!r} lacks a poll column: {err}") from err
    geography = next(
        (i for i, name in enumerate(names) if name in ("Georgraphy", "Geography")),
        None,
    )
    # Vote shares come first; later blocks (leader ratings, seat projections,
    # "wisdom of crowds") reuse the party names.
    end = names.index(RATINGS_HEADER) if RATINGS_HEADER in names else len(names)
    parties: dict[str, int] = {}
    for index, name in enumerate(names[:end]):
        if name is not None and name in PARTY_COLUMNS:
            key = PARTY_COLUMNS[name][0]
            if key in parties:
                raise PollBaseError(f"sheet {sheet!r} has two {key} columns")
            parties[key] = index
    missing = [key for key in REQUIRED_PARTIES if key not in parties]
    if missing:
        raise PollBaseError(f"sheet {sheet!r} lacks columns for {missing}")
    return _Columns(
        fieldwork=fieldwork,
        # The start date sits in the unlabelled column after the text.
        start=fieldwork + 1,
        published=published,
        polling=polling,
        publisher=publisher,
        geography=geography,
        parties=parties,
    )


def _cell(row: Sequence[object], index: int | None) -> object:
    if index is None or index >= len(row):
        return None
    return row[index]


def _rows_until_padding(rows: Iterable[Sequence[object]]) -> Iterator[Sequence[object]]:
    empty_run = 0
    for row in rows:
        if all(cell is None for cell in row):
            empty_run += 1
            if empty_run >= MAX_EMPTY_RUN:
                return
        else:
            empty_run = 0
        yield row


def read_workbook(
    path: Path,
    polling_days: Sequence[date],
    skipped_rows: Mapping[tuple[str, int, str, str], str] | None = None,
) -> Workbook:
    """Read the poll rows near each election, and every ``Result`` row.

    Only polls whose fieldwork started within ``READ_LOOKBACK`` before a
    polling day are parsed, so a quirk in a decades-old cell this builder
    never uses cannot stop it.

    Args:
        path: The ``.xlsx`` file.
        polling_days: The elections whose polls are read.
        skipped_rows: Rows to skip by (sheet, row, pollster, fieldwork
            text); :data:`SKIPPED_ROWS` when None.

    Returns:
        The rows of every sheet whose header has ``Fieldwork`` and
        ``Polling`` columns.

    Raises:
        PollBaseError: The file is not a readable workbook, a poll sheet
            lacks a column, a cell this builder reads cannot be parsed, a
            row's publication date falls before its fieldwork end, or a
            skipped row was not found.
    """
    skips = SKIPPED_ROWS if skipped_rows is None else skipped_rows
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except _CORRUPT_WORKBOOK as err:
        raise PollBaseError(f"{path} is not a readable workbook: {err}") from err
    try:
        workbook = Workbook()
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            columns = _columns(sheet.title, header)
            if columns is None:
                continue
            _read_sheet(sheet.title, columns, rows, polling_days, workbook, skips)
    except PollBaseError:
        raise
    except _CORRUPT_WORKBOOK as err:
        # read_only mode decompresses and parses each sheet as it is iterated.
        raise PollBaseError(f"{path} is not a readable workbook: {err}") from err
    finally:
        book.close()
    missing = sorted(set(skips) - set(workbook.skipped))
    if missing:
        raise PollBaseError(f"skipped rows not found (renumbered?): {missing}")
    return workbook


def _read_sheet(
    title: str,
    columns: _Columns,
    rows: Iterable[Sequence[object]],
    polling_days: Sequence[date],
    workbook: Workbook,
    skipped_rows: Mapping[tuple[str, int, str, str], str],
) -> None:
    last_start: date | None = None
    for number, row in enumerate(_rows_until_padding(rows), start=2):
        shares = {
            key: value
            for key, index in columns.parties.items()
            if (value := _number(_cell(row, index))) is not None
        }
        if _text(_cell(row, columns.polling)) == RESULT_LABEL:
            day = _number(_cell(row, columns.published))
            workbook.results.append(
                ResultRow(
                    sheet=title,
                    row=number,
                    day=None if day is None else int(day),
                    last_poll_start=last_start,
                    shares=shares,
                )
            )
            continue
        start_cell = _cell(row, columns.start)
        if not isinstance(start_cell, datetime):
            continue
        start = start_cell.date()
        last_start = start if last_start is None else max(last_start, start)
        if not any(day - READ_LOOKBACK <= start <= day for day in polling_days):
            continue
        pollster = _text(_cell(row, columns.polling))
        if pollster is None or any(key not in shares for key in REQUIRED_PARTIES):
            continue
        fieldwork_cell = _cell(row, columns.fieldwork)
        key = (title, number, pollster, _text(fieldwork_cell) or "")
        if key in skipped_rows:
            workbook.skipped[key] = skipped_rows[key]
            continue
        try:
            end = fieldwork_end(fieldwork_cell, start)
            published = published_date(_cell(row, columns.published), start)
            if published is not None and published < end:
                # A day number read forward from the start that lands before
                # the end is in the wrong month (or the cell is wrong).
                raise PollBaseError(
                    f"published {published} before the fieldwork end {end}"
                )
        except PollBaseError as err:
            raise PollBaseError(f"sheet {title!r} row {number}: {err}") from err
        workbook.polls.append(
            PollRow(
                sheet=title,
                row=number,
                pollster=pollster,
                publisher=_text(_cell(row, columns.publisher)),
                fieldwork=_text(fieldwork_cell),
                start=start,
                end=end,
                published=published,
                geography=_geography(_cell(row, columns.geography)),
                shares=shares,
            )
        )


# ── Selecting the final polls ─────────────────────────────────────────────────


@dataclass(frozen=True, slots=True)
class Exclusion:
    """A poll in the final-days window that the rule leaves out.

    Attributes:
        poll: The row.
        reason: Why it is not one of the final polls.
    """

    poll: PollRow
    reason: str


@dataclass(frozen=True, slots=True)
class ElectionPolls:
    """One election's final polls, average and result.

    Attributes:
        polling_day: The election's date.
        window_start: First fieldwork end date the window accepts.
        final_polls: Each pollster's final poll, ordered by pollster.
        excluded: Window polls left out, with the reason.
        result: The election's ``Result`` row.
        parties: Party keys averaged (reported by the result and at least
            one final poll), in ``PARTY_COLUMNS`` order.
        averages: Party key -> equal-weight mean over the reporting
            pollsters.
        reporting: Party key -> number of pollsters reporting it.
    """

    polling_day: date
    window_start: date
    final_polls: tuple[PollRow, ...]
    excluded: tuple[Exclusion, ...]
    result: ResultRow
    parties: tuple[str, ...]
    averages: Mapping[str, float]
    reporting: Mapping[str, int]


def _party_order() -> list[str]:
    order: list[str] = []
    for key, _label in PARTY_COLUMNS.values():
        if key not in order:
            order.append(key)
    return order


def find_result(results: Sequence[ResultRow], polling_day: date) -> ResultRow:
    """Return the ``Result`` row of the election held on *polling_day*.

    Args:
        results: Every result row in the workbook.
        polling_day: The election's date.

    Returns:
        The one row whose day matches and whose sheet's polls end within
        ``RESULT_LOOKBACK`` before polling day.

    Raises:
        PollBaseError: No row, or more than one, matches.
    """
    matches = [
        result
        for result in results
        if result.day == polling_day.day
        and result.last_poll_start is not None
        and polling_day - RESULT_LOOKBACK <= result.last_poll_start <= polling_day
    ]
    if len(matches) != 1:
        raise PollBaseError(
            f"expected one Result row for {polling_day}, found {len(matches)}"
        )
    return matches[0]


def _final_key(poll: PollRow) -> tuple[date, date, int]:
    return (poll.end, poll.published or date.min, poll.row)


def select_final_polls(
    polls: Iterable[PollRow],
    polling_day: date,
    *,
    window_days: int = WINDOW_DAYS,
) -> tuple[list[PollRow], list[Exclusion]]:
    """Apply the module's final-poll rule to one election.

    Args:
        polls: Every poll row in the workbook.
        polling_day: The election's date.
        window_days: Days before polling day the fieldwork may end on.

    Returns:
        Each pollster's final poll (sorted by pollster), and the window's
        other polls with the reason each was left out.
    """
    window_start = polling_day - timedelta(days=window_days)
    in_window = [
        poll
        for poll in polls
        if window_start <= poll.end <= polling_day and poll.start < polling_day
    ]
    excluded: list[Exclusion] = []
    qualifying: dict[str, list[PollRow]] = {}
    for poll in in_window:
        if poll.published is None:
            excluded.append(Exclusion(poll, "no publication date"))
        elif poll.published > polling_day:
            excluded.append(
                Exclusion(poll, f"published after polling day ({poll.published})")
            )
        else:
            qualifying.setdefault(poll.pollster, []).append(poll)
    finals: list[PollRow] = []
    for pollster in sorted(qualifying):
        ranked = sorted(qualifying[pollster], key=_final_key)
        finals.append(ranked[-1])
        excluded.extend(
            Exclusion(poll, "superseded by a later poll from the same pollster")
            for poll in ranked[:-1]
        )
    excluded.sort(key=lambda item: (item.poll.sheet, item.poll.row))
    return finals, excluded


def build_election(workbook: Workbook, polling_day: date) -> ElectionPolls:
    """Select one election's final polls and average them.

    Args:
        workbook: The rows read by :func:`read_workbook`.
        polling_day: The election's date.

    Returns:
        The election's final polls, averages and result.

    Raises:
        PollBaseError: There is no result row, no final poll, or the result
            lacks a required party.
    """
    result = find_result(workbook.results, polling_day)
    missing = [key for key in REQUIRED_PARTIES if key not in result.shares]
    if missing:
        raise PollBaseError(f"{polling_day} result lacks {missing}")
    finals, excluded = select_final_polls(workbook.polls, polling_day)
    if not finals:
        raise PollBaseError(f"no final polls for {polling_day}")
    parties: list[str] = []
    averages: dict[str, float] = {}
    reporting: dict[str, int] = {}
    for key in _party_order():
        values = [poll.shares[key] for poll in finals if key in poll.shares]
        if key not in result.shares or not values:
            continue
        parties.append(key)
        averages[key] = fmean(values)
        reporting[key] = len(values)
    return ElectionPolls(
        polling_day=polling_day,
        window_start=polling_day - timedelta(days=WINDOW_DAYS),
        final_polls=tuple(finals),
        excluded=tuple(excluded),
        result=result,
        parties=tuple(parties),
        averages=averages,
        reporting=reporting,
    )


# ── Output ────────────────────────────────────────────────────────────────────


def _labels() -> dict[str, str]:
    labels: dict[str, str] = {}
    for key, label in PARTY_COLUMNS.values():
        labels.setdefault(key, label)
    return labels


def _poll_json(poll: PollRow, parties: Sequence[str]) -> dict[str, object]:
    return {
        "pollster": poll.pollster,
        "publisher": poll.publisher,
        "fieldwork": poll.fieldwork,
        "fieldwork_start": poll.start.isoformat(),
        "fieldwork_end": poll.end.isoformat(),
        "published": None if poll.published is None else poll.published.isoformat(),
        "geography": poll.geography,
        "sheet": poll.sheet,
        "row": poll.row,
        "shares": {key: poll.shares[key] for key in parties if key in poll.shares},
    }


def election_json(election: ElectionPolls, edition: str) -> dict[str, object]:
    """Return one election's entry of the curated table.

    Args:
        election: The built election.
        edition: The PollBase file name, cited in the source.

    Returns:
        A JSON-ready dict.
    """
    labels = _labels()
    result = election.result
    return {
        "year": election.polling_day.year,
        "polling_day": election.polling_day.isoformat(),
        "window": [
            election.window_start.isoformat(),
            election.polling_day.isoformat(),
        ],
        "source": {
            "name": POLLBASE_CREDIT,
            "url": POLLBASE_URL,
            "edition": edition,
            "result": f"sheet {result.sheet!r} row {result.row} ({RESULT_LABEL})",
        },
        "parties": {
            key: {
                "label": labels[key],
                "average": round(election.averages[key], 4),
                "result": result.shares[key],
                "pollsters": election.reporting[key],
            }
            for key in election.parties
        },
        "final_polls": [
            _poll_json(poll, election.parties) for poll in election.final_polls
        ],
        "excluded_polls": [
            {**_poll_json(item.poll, election.parties), "reason": item.reason}
            for item in election.excluded
        ],
    }


def build_elections(
    workbook: Workbook, polling_days: Sequence[date]
) -> list[ElectionPolls]:
    """Build every election of *polling_days* from the workbook's rows."""
    return [build_election(workbook, day) for day in polling_days]


def table_json(
    elections: Sequence[ElectionPolls],
    edition: str,
    digest: str,
    skipped: Mapping[tuple[str, int, str, str], str] | None = None,
) -> dict[str, object]:
    """Return the whole curated table.

    Args:
        elections: The built elections.
        edition: The PollBase file name.
        digest: The file's SHA-256, so the edition used can be checked.
        skipped: The rows the read skipped (``Workbook.skipped``), listed
            with their reasons.

    Returns:
        The JSON-ready table.
    """
    return {
        "schema": FINAL_POLL_SCHEMA,
        "credit": {
            "text": f"Final-poll data: {POLLBASE_CREDIT} — {POLLBASE_URL}",
            "name": POLLBASE_CREDIT,
            "url": POLLBASE_URL,
            "edition": edition,
            "sha256": digest,
        },
        "rule": {
            "window_days": WINDOW_DAYS,
            "text": RULE_TEXT,
        },
        "skipped_rows": [
            {
                "sheet": sheet,
                "row": row,
                "pollster": pollster,
                "fieldwork": fieldwork,
                "reason": reason,
            }
            for (sheet, row, pollster, fieldwork), reason in (skipped or {}).items()
        ],
        "elections": [election_json(election, edition) for election in elections],
    }


def summary_lines(elections: Sequence[ElectionPolls]) -> list[str]:
    """Return a printable summary: per election, A vs R per party."""
    lines: list[str] = []
    for election in elections:
        lines.append(
            f"{election.polling_day}: {len(election.final_polls)} final polls, "
            f"{len(election.excluded)} excluded"
        )
        for poll in election.final_polls:
            shares = " ".join(
                f"{key}={poll.shares[key]:g}"
                for key in election.parties
                if key in poll.shares
            )
            lines.append(
                f"    {poll.pollster:<16} {poll.fieldwork or '-':<8} "
                f"end {poll.end} pub {poll.published} "
                f"{poll.geography or '':<3} {shares}"
            )
        for item in election.excluded:
            lines.append(
                f"    x {item.poll.pollster:<14} {item.poll.fieldwork or '-':<8} "
                f"end {item.poll.end}: {item.reason}"
            )
        for key in election.parties:
            lines.append(
                f"  {key:<7} A {election.averages[key]:6.2f}  "
                f"R {election.result.shares[key]:6.2f}  "
                f"({election.reporting[key]} pollsters)"
            )
    return lines


def main(argv: Sequence[str] | None = None) -> int:
    """Build the curated table and write it (unless ``--dry-run``).

    Args:
        argv: Command-line arguments; defaults to ``sys.argv[1:]``.

    Returns:
        Process exit code: 0 on success, 1 if the workbook cannot be read.
    """
    parser = argparse.ArgumentParser(
        description="Build uk_final_poll_errors.json from a PollBase workbook.",
    )
    parser.add_argument("workbook", type=Path, help="The PollBase .xlsx file.")
    parser.add_argument(
        "--output",
        type=Path,
        default=FINAL_POLL_ERRORS_PATH,
        help="Where to write the table.",
    )
    parser.add_argument(
        "--dry-run", action="store_true", help="Print the table without writing it."
    )
    args = parser.parse_args(argv)
    try:
        workbook = read_workbook(args.workbook, POLLING_DAYS)
        elections = build_elections(workbook, POLLING_DAYS)
        digest = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    except (PollBaseError, OSError) as err:
        print(f"error: {err}")
        return 1
    for line in summary_lines(elections):
        print(line)
    if args.dry_run:
        print("\nDry-run complete. Nothing written.")
        return 0
    table = table_json(elections, args.workbook.name, digest, workbook.skipped)
    text = json.dumps(table, indent=2, ensure_ascii=False) + "\n"
    try:
        write_text_files({args.output: text})
    except OSError as err:
        print(f"error: {err}")
        return 1
    print(f"\nWrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
