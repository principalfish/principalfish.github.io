"""Tests for the BMG Research XLSX poll importer.

Covers parsing helpers, ``build_import_plan``, ``_cli_preview`` and ``main``.
``commit_import_plan`` and ``_find_existing_poll`` are already covered for every
Westminster importer, including this one, by
``test_westminster_importers_commit.py`` and are not retested here.
"""

from __future__ import annotations

import sys
from collections.abc import Mapping
from datetime import date
from types import MappingProxyType
from typing import Any
from urllib.request import Request

import pytest
from openpyxl import Workbook

from db import Database
from polls.importers.westminster import bmg_research_import as bmg
from tests.uk_fixtures import (
    FakeUrlResponse,
    WestminsterWorld,
    build_workbook,
    workbook_bytes,
)

# ── Synthetic workbook builders ────────────────────────────────────────────────
#
# The "Tables" sheet layout the parser expects: a "WouldVoteTodayRevised" marker
# row, then a national block (a party-name row in column B, followed by its
# percentage in column C one row below); a second marker row, two blank rows, a
# region-header row (region names starting at column 34), then a regional block
# (a party-name row, followed by its per-region percentages one row below).

_COL_WIDTH = 90
_VI_MARKER = "WouldVoteTodayRevised"

_REQUIRED_PARTIES: tuple[str, ...] = (
    "Conservative",
    "Labour",
    "Liberal Democrats",
    "Reform UK",
    "Green",
    "Scottish National Party",
    "Plaid Cymru",
    "Other",
)

# Fractions in [0.0, 1.0]; _to_percentage multiplies by 100. Sums to 1.0.
_NATIONAL_FRACTIONS: Mapping[str, float] = MappingProxyType(
    {
        "Conservative": 0.24,
        "Labour": 0.32,
        "Liberal Democrats": 0.11,
        "Reform UK": 0.21,
        "Green": 0.06,
        "Scottish National Party": 0.03,
        "Plaid Cymru": 0.01,
        "Other": 0.02,
    }
)

# Source column headers, spelled as BMG spells them, mapped to sheet columns.
_REGION_COLUMNS: Mapping[str, int] = MappingProxyType(
    {"London": 34, "Scotland": 35, "North East": 36}
)

# Only Labour and Conservative get regional rows, so every other party's
# regions come out at the 0.0 default.
_REGIONAL_FRACTIONS: Mapping[str, Mapping[str, float]] = MappingProxyType(
    {
        "Labour": MappingProxyType(
            {"London": 0.42, "Scotland": 0.30, "North East": 0.33}
        ),
        "Conservative": MappingProxyType(
            {"London": 0.18, "Scotland": 0.16, "North East": 0.19}
        ),
    }
)

_FIELDWORK_TEXT = "Fieldwork dates: 3-4 January 2026"
_SAMPLE_TEXT = "Sample: 1500 GB adults"
_XLSX_URL = "https://bmgresearch.com/wp-content/uploads/2026/02/tables.xlsx"


def _row(values: Mapping[int, object]) -> list[object]:
    """One 1-indexed worksheet row (as a 0-indexed list), padded to _COL_WIDTH."""
    cells: list[object] = [None] * _COL_WIDTH
    for col, value in values.items():
        cells[col - 1] = value
    return cells


def _national_block(percentages: Mapping[str, float]) -> list[list[object]]:
    """A label row (col B) plus a value row (col C, the row below) per party."""
    rows: list[list[object]] = []
    for party, fraction in percentages.items():
        rows.append(_row({2: party}))
        rows.append(_row({3: fraction}))
    return rows


def _regional_block(
    region_columns: Mapping[str, int], percentages: Mapping[str, Mapping[str, float]]
) -> list[list[object]]:
    """A label row (col B) plus a value row (region columns) per party."""
    rows: list[list[object]] = []
    for party, region_fractions in percentages.items():
        rows.append(_row({2: party}))
        values = {
            region_columns[region]: fraction
            for region, fraction in region_fractions.items()
        }
        rows.append(_row(values))
    return rows


def _wrap_with_regional_trailer(
    national_rows: list[list[object]],
    *,
    region_columns: Mapping[str, int] = _REGION_COLUMNS,
    regional: Mapping[str, Mapping[str, float]] = _REGIONAL_FRACTIONS,
) -> list[list[object]]:
    """Append the marker/blanks/header/regional-block trailer after national rows."""
    rows = list(national_rows)
    rows.append(_row({2: _VI_MARKER}))
    rows.append(_row({}))
    rows.append(_row({}))
    rows.append(_row({col: name for name, col in region_columns.items()}))
    rows.extend(_regional_block(region_columns, regional))
    return rows


def _tables_sheet(
    *,
    national: Mapping[str, float] = _NATIONAL_FRACTIONS,
    region_columns: Mapping[str, int] = _REGION_COLUMNS,
    regional: Mapping[str, Mapping[str, float]] = _REGIONAL_FRACTIONS,
) -> list[list[object]]:
    """A "Tables" sheet: a national VI block, then a regional VI block."""
    rows: list[list[object]] = [_row({2: _VI_MARKER})]
    rows.extend(_national_block(national))
    rows.append(_row({2: _VI_MARKER}))
    rows.append(_row({}))
    rows.append(_row({}))
    rows.append(_row({col: name for name, col in region_columns.items()}))
    rows.extend(_regional_block(region_columns, regional))
    return rows


def _methodology_sheet(
    *,
    fieldwork_text: str | None = _FIELDWORK_TEXT,
    sample_text: str | None = _SAMPLE_TEXT,
) -> list[list[object]]:
    """A "Methodology" sheet with a fieldwork-dates line and a sample line."""
    rows: list[list[object]] = [["BMG Research Omnibus"]]
    if fieldwork_text is not None:
        rows.append([fieldwork_text])
    if sample_text is not None:
        rows.append([sample_text])
    return rows


def _full_workbook(
    *,
    fieldwork_text: str | None = _FIELDWORK_TEXT,
    sample_text: str | None = _SAMPLE_TEXT,
    national: Mapping[str, float] = _NATIONAL_FRACTIONS,
    region_columns: Mapping[str, int] = _REGION_COLUMNS,
    regional: Mapping[str, Mapping[str, float]] = _REGIONAL_FRACTIONS,
) -> Workbook:
    """A full workbook: a "Methodology" sheet plus a "Tables" sheet."""
    return build_workbook(
        {
            "Methodology": _methodology_sheet(
                fieldwork_text=fieldwork_text, sample_text=sample_text
            ),
            "Tables": _tables_sheet(
                national=national, region_columns=region_columns, regional=regional
            ),
        }
    )


# ── _month_number ────────────────────────────────────────────────────────────


class TestMonthNumber:
    """Tests for _month_number — month name → integer conversion."""

    def test_full_names(self) -> None:
        assert bmg._month_number("January") == 1
        assert bmg._month_number("February") == 2
        assert bmg._month_number("March") == 3
        assert bmg._month_number("April") == 4
        assert bmg._month_number("May") == 5
        assert bmg._month_number("June") == 6
        assert bmg._month_number("July") == 7
        assert bmg._month_number("August") == 8
        assert bmg._month_number("September") == 9
        assert bmg._month_number("October") == 10
        assert bmg._month_number("November") == 11
        assert bmg._month_number("December") == 12

    def test_abbreviated_names(self) -> None:
        assert bmg._month_number("Jan") == 1
        assert bmg._month_number("Feb") == 2
        assert bmg._month_number("Sep") == 9
        assert bmg._month_number("Sept") == 9
        assert bmg._month_number("Dec") == 12

    def test_case_insensitive(self) -> None:
        assert bmg._month_number("JANUARY") == 1
        assert bmg._month_number("january") == 1
        assert bmg._month_number("jAnUaRy") == 1

    def test_trailing_period_stripped(self) -> None:
        assert bmg._month_number("Jan.") == 1

    def test_unknown_returns_none(self) -> None:
        assert bmg._month_number("Octember") is None
        assert bmg._month_number("") is None
        assert bmg._month_number("13") is None


# ── _infer_year ───────────────────────────────────────────────────────────────


class TestInferYear:
    """Tests for _infer_year — four-digit year extraction from URLs."""

    def test_year_as_path_segment(self) -> None:
        url = "https://bmgresearch.com/wp-content/uploads/2026/02/tables.xlsx"
        assert bmg._infer_year(url) == 2026

    def test_year_as_bare_occurrence(self) -> None:
        url = "https://bmgresearch.com/tables_2025_v2.xlsx"
        assert bmg._infer_year(url) == 2025

    def test_path_segment_preferred_over_bare(self) -> None:
        url = "https://bmgresearch.com/2024archive/2026/tables.xlsx"
        assert bmg._infer_year(url) == 2026

    def test_fallback_when_no_year_in_url(self) -> None:
        url = "https://bmgresearch.com/tables.xlsx"
        assert bmg._infer_year(url, fallback=2025) == 2025

    def test_no_year_no_fallback_raises(self) -> None:
        with pytest.raises(ValueError, match="Could not infer year"):
            bmg._infer_year("https://bmgresearch.com/tables.xlsx")


# ── _parse_fieldwork ──────────────────────────────────────────────────────────


class TestParseFieldwork:
    """Tests for _parse_fieldwork — date-range string parsing."""

    def test_same_month_ampersand_with_year(self) -> None:
        start, end = bmg._parse_fieldwork("3 & 4 January 2026")
        assert start == date(2026, 1, 3)
        assert end == date(2026, 1, 4)

    def test_same_month_hyphen_with_year(self) -> None:
        start, end = bmg._parse_fieldwork("3-4 January 2026")
        assert start == date(2026, 1, 3)
        assert end == date(2026, 1, 4)

    def test_cross_month_with_year(self) -> None:
        start, end = bmg._parse_fieldwork("31 January - 2 February 2026")
        assert start == date(2026, 1, 31)
        assert end == date(2026, 2, 2)

    def test_cross_month_year_wraps_backward(self) -> None:
        start, end = bmg._parse_fieldwork("31 December - 2 January 2026")
        assert start == date(2025, 12, 31)
        assert end == date(2026, 1, 2)

    def test_no_year_range_uses_default(self) -> None:
        start, end = bmg._parse_fieldwork("3-4 January", default_year=2026)
        assert start == date(2026, 1, 3)
        assert end == date(2026, 1, 4)

    def test_single_day_with_year(self) -> None:
        start, end = bmg._parse_fieldwork("3 January 2026")
        assert start == date(2026, 1, 3)
        assert end == date(2026, 1, 3)

    def test_single_day_no_year_uses_default(self) -> None:
        start, end = bmg._parse_fieldwork("3 January", default_year=2026)
        assert start == date(2026, 1, 3)
        assert end == date(2026, 1, 3)

    def test_ordinal_suffixes_stripped(self) -> None:
        start, end = bmg._parse_fieldwork("3rd & 4th January 2026")
        assert start == date(2026, 1, 3)
        assert end == date(2026, 1, 4)

    @pytest.mark.parametrize(
        "dash", ["–", "—"], ids=["en_dash", "em_dash"]
    )
    def test_dash_variants_normalised(self, dash: str) -> None:
        start, end = bmg._parse_fieldwork(f"31 January {dash} 2 February 2026")
        assert start == date(2026, 1, 31)
        assert end == date(2026, 2, 2)

    def test_no_pattern_matches_raises(self) -> None:
        with pytest.raises(ValueError, match="Could not parse fieldwork string"):
            bmg._parse_fieldwork("not a date")

    def test_no_year_pattern_without_default_year_falls_through(self) -> None:
        # "3-4 January" without default_year matches no pattern at all: the
        # no-year pattern requires default_year, and nothing else fits.
        with pytest.raises(ValueError, match="Could not parse fieldwork string"):
            bmg._parse_fieldwork("3-4 January")

    @pytest.mark.parametrize(
        "text",
        [
            "3 & 4 Blorpuary 2026",
            "3-4 Blorpuary 2026",
            "31 Blorpuary - 2 February 2026",
            "3 Blorpuary 2026",
        ],
        ids=["same_month_ampersand", "same_month_hyphen", "cross_month", "single_day"],
    )
    def test_unrecognised_month_raises(self, text: str) -> None:
        with pytest.raises(ValueError, match="Could not parse fieldwork month"):
            bmg._parse_fieldwork(text)

    def test_no_year_pattern_unrecognised_month_raises(self) -> None:
        with pytest.raises(ValueError, match="Could not parse fieldwork month"):
            bmg._parse_fieldwork("3-4 Blorpuary", default_year=2026)

    def test_single_day_no_year_pattern_unrecognised_month_raises(self) -> None:
        with pytest.raises(ValueError, match="Could not parse fieldwork month"):
            bmg._parse_fieldwork("3 Blorpuary", default_year=2026)


# ── extract_workbook ────────────────────────────────────────────────────────


def _valid_payload() -> bytes:
    """XLSX bytes for a trivial one-sheet workbook (starts with the PK magic)."""
    return workbook_bytes(build_workbook({"Sheet1": [["ok"]]}))


class TestExtractWorkbook:
    """Tests for extract_workbook — fetch and load an XLSX from a URL."""

    def test_plain_url_success(self, monkeypatch: pytest.MonkeyPatch) -> None:
        payload = _valid_payload()

        def fake_urlopen(req: Request, timeout: float = 45) -> FakeUrlResponse:
            assert req.full_url == "https://bmgresearch.com/tables.xlsx"
            assert req.get_header("User-agent") == "Mozilla/5.0"
            assert req.get_header("Referer") == "https://bmgresearch.com/"
            return FakeUrlResponse(payload)

        monkeypatch.setattr(bmg, "urlopen", fake_urlopen)
        workbook = bmg.extract_workbook("https://bmgresearch.com/tables.xlsx")
        assert workbook.sheetnames == ["Sheet1"]

    def test_non_xlsx_payload_raises(self, monkeypatch: pytest.MonkeyPatch) -> None:
        def fake_urlopen(req: Request, timeout: float = 45) -> FakeUrlResponse:
            return FakeUrlResponse(b"<html>not an xlsx file</html>")

        monkeypatch.setattr(bmg, "urlopen", fake_urlopen)
        with pytest.raises(ValueError, match="non-xlsx payload"):
            bmg.extract_workbook("https://bmgresearch.com/tables.xlsx")

    def test_urlopen_exception_is_captured_in_message(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        def fake_urlopen(req: Request, timeout: float = 45) -> FakeUrlResponse:
            raise TimeoutError("connection timed out")

        monkeypatch.setattr(bmg, "urlopen", fake_urlopen)
        with pytest.raises(ValueError, match="connection timed out"):
            bmg.extract_workbook("https://bmgresearch.com/tables.xlsx")

    def test_co_uk_url_falls_back_through_com_variants(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        payload = _valid_payload()
        requested: list[str] = []

        def fake_urlopen(req: Request, timeout: float = 45) -> FakeUrlResponse:
            requested.append(req.full_url)
            if req.full_url == "https://bmgresearch.com/foo.xlsx":
                return FakeUrlResponse(payload)
            raise ConnectionError("unreachable")

        monkeypatch.setattr(bmg, "urlopen", fake_urlopen)
        workbook = bmg.extract_workbook("https://www.bmgresearch.co.uk/foo.xlsx")

        assert workbook.sheetnames == ["Sheet1"]
        assert requested == [
            "https://www.bmgresearch.co.uk/foo.xlsx",
            "https://www.bmgresearch.com/foo.xlsx",
            "https://bmgresearch.com/foo.xlsx",
        ]

    def test_co_uk_without_www_tries_exactly_two_candidates(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        # Without a "www." prefix, the second replacement (of
        # "www.bmgresearch.co.uk") matches nothing and leaves the original URL
        # unchanged, so dict.fromkeys should dedupe it away rather than
        # requesting it (or the .com variant) twice.
        payload = _valid_payload()
        requested: list[str] = []

        def fake_urlopen(req: Request, timeout: float = 45) -> FakeUrlResponse:
            requested.append(req.full_url)
            if req.full_url == "https://bmgresearch.com/foo.xlsx":
                return FakeUrlResponse(payload)
            raise ConnectionError("unreachable")

        monkeypatch.setattr(bmg, "urlopen", fake_urlopen)
        workbook = bmg.extract_workbook("https://bmgresearch.co.uk/foo.xlsx")

        assert workbook.sheetnames == ["Sheet1"]
        assert requested == [
            "https://bmgresearch.co.uk/foo.xlsx",
            "https://bmgresearch.com/foo.xlsx",
        ]

    def test_all_candidates_failing_combines_errors(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        def fake_urlopen(req: Request, timeout: float = 45) -> FakeUrlResponse:
            raise ConnectionError("unreachable")

        monkeypatch.setattr(bmg, "urlopen", fake_urlopen)
        with pytest.raises(
            ValueError, match="Could not fetch XLSX payload"
        ) as exc_info:
            bmg.extract_workbook("https://www.bmgresearch.co.uk/foo.xlsx")
        message = str(exc_info.value)
        assert message.count("unreachable") == 3
        assert " | " in message


# ── _cell_text ────────────────────────────────────────────────────────────────


class TestCellText:
    """Tests for _cell_text — openpyxl cell value → stripped string."""

    def test_none_returns_empty_string(self) -> None:
        assert bmg._cell_text(None) == ""

    def test_string_stripped(self) -> None:
        assert bmg._cell_text("  hello  ") == "hello"

    def test_integer_converted(self) -> None:
        assert bmg._cell_text(42) == "42"

    def test_float_converted(self) -> None:
        assert bmg._cell_text(3.14) == "3.14"

    def test_empty_string(self) -> None:
        assert bmg._cell_text("") == ""


# ── _to_percentage / _to_percentage_or_zero ────────────────────────────────────


class TestToPercentage:
    """Tests for _to_percentage — raw cell value → percentage float."""

    def test_integer_in_0_100_range(self) -> None:
        assert bmg._to_percentage(42) == pytest.approx(42.0)

    def test_decimal_in_0_1_range_multiplied(self) -> None:
        assert bmg._to_percentage(0.42) == pytest.approx(42.0)

    def test_rounds_to_nearest_integer(self) -> None:
        assert bmg._to_percentage(34.6) == pytest.approx(35.0)
        assert bmg._to_percentage(34.4) == pytest.approx(34.0)

    def test_zero(self) -> None:
        assert bmg._to_percentage(0) == pytest.approx(0.0)

    def test_one_boundary_treated_as_100_percent(self) -> None:
        assert bmg._to_percentage(1.0) == pytest.approx(100.0)

    def test_returns_a_float(self) -> None:
        assert type(bmg._to_percentage(42)) is float

    def test_none_raises(self) -> None:
        with pytest.raises(ValueError, match="empty percentage cell"):
            bmg._to_percentage(None)


class TestToPercentageOrZero:
    """Tests for _to_percentage_or_zero — blank-tolerant percentage conversion."""

    def test_none_returns_zero(self) -> None:
        assert bmg._to_percentage_or_zero(None) == pytest.approx(0.0)

    def test_hyphen_returns_zero(self) -> None:
        assert bmg._to_percentage_or_zero("-") == pytest.approx(0.0)

    def test_en_dash_returns_zero(self) -> None:
        assert bmg._to_percentage_or_zero("–") == pytest.approx(0.0)

    def test_em_dash_returns_zero(self) -> None:
        assert bmg._to_percentage_or_zero("—") == pytest.approx(0.0)

    def test_empty_string_returns_zero(self) -> None:
        assert bmg._to_percentage_or_zero("") == pytest.approx(0.0)

    def test_numeric_delegates_to_to_percentage(self) -> None:
        assert bmg._to_percentage_or_zero(0.35) == pytest.approx(35.0)

    def test_non_dash_string_delegates_to_to_percentage(self) -> None:
        assert bmg._to_percentage_or_zero("42") == pytest.approx(42.0)

    def test_integer_in_0_100_range(self) -> None:
        assert bmg._to_percentage_or_zero(28) == pytest.approx(28.0)


# ── _normalize_header ──────────────────────────────────────────────────────────


class TestNormalizeHeader:
    """Tests for _normalize_header — whitespace-collapsing header cleanup."""

    def test_collapses_internal_whitespace(self) -> None:
        assert bmg._normalize_header("East   Midlands") == "East Midlands"

    def test_strips_leading_and_trailing_whitespace(self) -> None:
        assert bmg._normalize_header("  Scotland  ") == "Scotland"

    def test_tabs_and_newlines_collapsed(self) -> None:
        assert bmg._normalize_header("North\t\nWest") == "North West"

    def test_empty_string(self) -> None:
        assert bmg._normalize_header("") == ""

    def test_already_normalised_unchanged(self) -> None:
        assert bmg._normalize_header("London") == "London"


# ── _find_methodology_sheet ─────────────────────────────────────────────────────


class TestFindMethodologySheet:
    """Tests for _find_methodology_sheet — locate the methodology sheet."""

    def test_matches_sheet_containing_method_case_insensitively(self) -> None:
        workbook = build_workbook({"Cover": [[1]], "METHODOLOGY notes": [[2]]})
        sheet = bmg._find_methodology_sheet(workbook)
        assert sheet.title == "METHODOLOGY notes"

    def test_falls_back_to_first_sheet_when_none_matches(self) -> None:
        workbook = build_workbook({"Cover": [[1]], "Data": [[2]]})
        sheet = bmg._find_methodology_sheet(workbook)
        assert sheet.title == "Cover"


# ── _parse_fieldwork_and_sample ─────────────────────────────────────────────────


class TestParseFieldworkAndSample:
    """Tests for _parse_fieldwork_and_sample — methodology sheet extraction."""

    def test_happy_path(self) -> None:
        workbook = build_workbook({"Methodology": _methodology_sheet()})
        start, end, sample = bmg._parse_fieldwork_and_sample(workbook, 2026)
        assert start == date(2026, 1, 3)
        assert end == date(2026, 1, 4)
        assert sample == 1500

    def test_missing_fieldwork_raises(self) -> None:
        workbook = build_workbook(
            {"Methodology": _methodology_sheet(fieldwork_text=None)}
        )
        with pytest.raises(ValueError, match="Fieldwork dates not found"):
            bmg._parse_fieldwork_and_sample(workbook, 2026)

    def test_missing_sample_raises(self) -> None:
        workbook = build_workbook({"Methodology": _methodology_sheet(sample_text=None)})
        with pytest.raises(ValueError, match="Sample line not found"):
            bmg._parse_fieldwork_and_sample(workbook, 2026)

    def test_sample_line_without_digits_raises(self) -> None:
        workbook = build_workbook(
            {"Methodology": _methodology_sheet(sample_text="Sample: undisclosed")}
        )
        with pytest.raises(ValueError, match="Could not parse sample size"):
            bmg._parse_fieldwork_and_sample(workbook, 2026)


# ── _find_tables_sheet ───────────────────────────────────────────────────────


class TestFindTablesSheet:
    """Tests for _find_tables_sheet — locate the tables sheet by name."""

    def test_matches_exact_name_tables_case_insensitively(self) -> None:
        workbook = build_workbook({"Cover": [[1]], "TABLES": [[2]]})
        assert bmg._find_tables_sheet(workbook).title == "TABLES"

    def test_matches_name_containing_table(self) -> None:
        workbook = build_workbook({"Cover": [[1]], "VI Tables Q1": [[2]]})
        assert bmg._find_tables_sheet(workbook).title == "VI Tables Q1"

    def test_missing_raises(self) -> None:
        workbook = build_workbook({"Cover": [[1]], "Data": [[2]]})
        with pytest.raises(ValueError, match="Could not find tables sheet"):
            bmg._find_tables_sheet(workbook)


# ── _find_vi_table_starts ────────────────────────────────────────────────────


class TestFindViTableStarts:
    """Tests for _find_vi_table_starts — locate WouldVoteTodayRevised headers."""

    def test_finds_every_occurrence_in_column_b(self) -> None:
        workbook = build_workbook({"Tables": _tables_sheet()})
        starts = bmg._find_vi_table_starts(workbook["Tables"])
        assert starts == [1, 18]

    def test_three_occurrences_returns_all_three(self) -> None:
        rows: list[list[object]] = [
            _row({2: _VI_MARKER}),
            _row({2: _VI_MARKER}),
            _row({2: _VI_MARKER}),
        ]
        workbook = build_workbook({"Tables": rows})
        starts = bmg._find_vi_table_starts(workbook["Tables"])
        assert starts == [1, 2, 3]

    def test_missing_raises(self) -> None:
        workbook = build_workbook({"Tables": [["nothing", "here"]]})
        with pytest.raises(ValueError, match="Could not find WouldVoteTodayRevised"):
            bmg._find_vi_table_starts(workbook["Tables"])


# ── _parse_party_region_percentages ─────────────────────────────────────────────


class TestParsePartyRegionPercentages:
    """Tests for _parse_party_region_percentages — national and regional VI."""

    def test_happy_path_national_and_regional_values(self) -> None:
        workbook = build_workbook({"Tables": _tables_sheet()})
        parsed = bmg._parse_party_region_percentages(workbook)

        assert len(parsed) == 8
        assert parsed["Labour"][bmg.NATIONAL_KEY] == 32.0
        assert parsed["Labour"]["London"] == 42.0
        assert parsed["Labour"]["Scotland"] == 30.0
        assert parsed["Labour"]["North East England"] == 33.0
        # Conservative is the last regional row at the scan's upper bound.
        assert parsed["Conservative"]["London"] == 18.0
        assert parsed["Conservative"]["Scotland"] == 16.0
        assert parsed["Conservative"]["North East England"] == 19.0

    def test_region_without_a_workbook_column_is_absent_not_defaulted(self) -> None:
        # _parse_party_region_percentages only returns the regions it found a
        # column for; build_import_plan is what fills in the rest of the DB's
        # regions at 0.0 (see TestBuildImportPlan).
        workbook = build_workbook({"Tables": _tables_sheet()})
        parsed = bmg._parse_party_region_percentages(workbook)
        assert "East Midlands" not in parsed["Labour"]

    def test_party_absent_from_regional_block_defaults_every_region_to_zero(
        self,
    ) -> None:
        workbook = build_workbook({"Tables": _tables_sheet()})
        parsed = bmg._parse_party_region_percentages(workbook)
        # Green has a national row but no row in the regional block.
        assert parsed["Green"]["London"] == 0.0
        assert parsed["Green"]["Scotland"] == 0.0

    def _national_rows_with_labour_compact(
        self, label_col3: object, value_row_col3: object
    ) -> list[list[object]]:
        """A national block where Labour's label row also carries a column-C
        value (``label_col3``, the "current row" reading), and the row below
        it carries ``value_row_col3`` (the "next row" reading). Every other
        required party uses the normal two-row (label, value-below) shape.
        """
        others = {
            name: pct for name, pct in _NATIONAL_FRACTIONS.items() if name != "Labour"
        }
        rows: list[list[object]] = [_row({2: _VI_MARKER})]
        rows.append(_row({2: "Labour", 3: label_col3}))
        rows.append(_row({3: value_row_col3}))
        rows.extend(_national_block(others))
        return _wrap_with_regional_trailer(rows)

    def test_next_row_takes_precedence_when_both_readings_are_valid(self) -> None:
        # Both the label row's own column C (0.35) and the row below it
        # (0.32) are valid readings; the row below must win, since it is
        # checked first in the source.
        rows = self._national_rows_with_labour_compact(0.35, 0.32)
        workbook = build_workbook({"Tables": rows})
        parsed = bmg._parse_party_region_percentages(workbook)
        assert parsed["Labour"][bmg.NATIONAL_KEY] == 32.0

    def test_row_below_out_of_range_falls_back_to_label_row(self) -> None:
        # 412 is numeric but well outside the accepted [0.0, 1.2] range, so
        # the parser must fall back to the label row's own column C (0.35)
        # rather than accepting it outright.
        rows = self._national_rows_with_labour_compact(0.35, 412)
        workbook = build_workbook({"Tables": rows})
        parsed = bmg._parse_party_region_percentages(workbook)
        assert parsed["Labour"][bmg.NATIONAL_KEY] == 35.0

    def test_neither_reading_valid_skips_party_and_raises(self) -> None:
        # Neither the label row nor the row below it carries a usable
        # column-C value, so Labour is skipped entirely — and since it is a
        # required party, the workbook fails the missing-party check.
        rows = self._national_rows_with_labour_compact(None, None)
        workbook = build_workbook({"Tables": rows})
        with pytest.raises(ValueError, match="Missing expected party rows") as exc_info:
            bmg._parse_party_region_percentages(workbook)
        assert "Labour" in str(exc_info.value)

    def test_table_header_within_the_first_four_rows_does_not_break_early(
        self,
    ) -> None:
        # A "Table 1: ..." title often sits right after the marker in real
        # BMG workbooks; the "row > start + 4" guard must not treat it as
        # the break trigger this early, or the whole national block would be
        # skipped.
        rows: list[list[object]] = [_row({2: _VI_MARKER})]
        rows.append(_row({2: "Table 1: Westminster Voting Intention"}))
        rows.extend(_national_block(_NATIONAL_FRACTIONS))
        rows = _wrap_with_regional_trailer(rows)

        workbook = build_workbook({"Tables": rows})
        parsed = bmg._parse_party_region_percentages(workbook)
        assert len(parsed) == 8

    def test_regional_table_header_within_first_four_rows_does_not_break_early(
        self,
    ) -> None:
        rows: list[list[object]] = [_row({2: _VI_MARKER})]
        rows.extend(_national_block(_NATIONAL_FRACTIONS))
        rows.append(_row({2: _VI_MARKER}))
        rows.append(_row({2: "Table 1: Regional Voting Intention"}))
        rows.append(_row({}))
        rows.append(_row({col: name for name, col in _REGION_COLUMNS.items()}))
        rows.extend(_regional_block(_REGION_COLUMNS, _REGIONAL_FRACTIONS))

        workbook = build_workbook({"Tables": rows})
        parsed = bmg._parse_party_region_percentages(workbook)
        assert parsed["Labour"]["London"] == 42.0

    def test_unmapped_region_header_is_ignored(self) -> None:
        header = {col: name for name, col in _REGION_COLUMNS.items()}
        header[37] = "Atlantis"
        rows: list[list[object]] = [_row({2: _VI_MARKER})]
        rows.extend(_national_block(_NATIONAL_FRACTIONS))
        rows.append(_row({2: _VI_MARKER}))
        rows.append(_row({}))
        rows.append(_row({}))
        rows.append(_row(header))
        rows.extend(_regional_block(_REGION_COLUMNS, _REGIONAL_FRACTIONS))

        workbook = build_workbook({"Tables": rows})
        parsed = bmg._parse_party_region_percentages(workbook)
        assert set(parsed["Labour"].keys()) == {
            bmg.NATIONAL_KEY,
            "London",
            "Scotland",
            "North East England",
        }
        assert parsed["Labour"]["London"] == 42.0

    def test_regional_table_header_stops_block_early(self) -> None:
        # A "Table 3: ..." row more than 4 rows past the regional marker halts
        # the regional scan, so a party listed after it keeps its national
        # figure but every region defaults to 0.0.
        rows: list[list[object]] = [_row({2: _VI_MARKER})]
        rows.extend(_national_block(_NATIONAL_FRACTIONS))
        rows.append(_row({2: _VI_MARKER}))
        rows.append(_row({}))
        rows.append(_row({}))
        rows.append(_row({col: name for name, col in _REGION_COLUMNS.items()}))
        rows.append(_row({2: "Labour"}))
        rows.append(_row({34: 0.42, 35: 0.30, 36: 0.33}))
        rows.append(_row({2: "Table 3: something else"}))
        conservative_only = {"Conservative": _REGIONAL_FRACTIONS["Conservative"]}
        rows.extend(_regional_block(_REGION_COLUMNS, conservative_only))

        workbook = build_workbook({"Tables": rows})
        parsed = bmg._parse_party_region_percentages(workbook)
        assert parsed["Labour"]["London"] == 42.0
        assert parsed["Conservative"]["London"] == 0.0

    def test_missing_regional_columns_raises(self) -> None:
        workbook = build_workbook(
            {"Tables": _tables_sheet(region_columns={}, regional={})}
        )
        with pytest.raises(ValueError, match="Could not locate BMG regional columns"):
            bmg._parse_party_region_percentages(workbook)

    def test_missing_non_optional_party_raises(self) -> None:
        national = {
            name: pct for name, pct in _NATIONAL_FRACTIONS.items() if name != "Green"
        }
        workbook = build_workbook({"Tables": _tables_sheet(national=national)})
        with pytest.raises(ValueError, match="Missing expected party rows") as exc_info:
            bmg._parse_party_region_percentages(workbook)
        assert "Green" in str(exc_info.value)

    @pytest.mark.parametrize("party_name", _REQUIRED_PARTIES[-3:])
    def test_missing_optional_party_still_raises_pins_current_behaviour(
        self, party_name: str
    ) -> None:
        # The docstring says SNP/Plaid Cymru/Other default to 0.0 when absent
        # from the workbook, but the required-party check compares against
        # national_values, not the defaulted parsed dict — so a workbook
        # missing one of these "optional" parties raises too, contrary to the
        # documented behaviour.
        national = {
            name: pct for name, pct in _NATIONAL_FRACTIONS.items() if name != party_name
        }
        workbook = build_workbook({"Tables": _tables_sheet(national=national)})
        with pytest.raises(ValueError, match="Missing expected party rows") as exc_info:
            bmg._parse_party_region_percentages(workbook)
        assert party_name in str(exc_info.value)

    def test_table_header_stops_national_block_early(self) -> None:
        # A "Table 2: ..." row more than 4 rows past the marker halts the
        # national scan, so the parties listed after it are never read.
        parties = list(_NATIONAL_FRACTIONS.items())
        before, after = dict(parties[:2]), dict(parties[2:])

        rows: list[list[object]] = [_row({2: _VI_MARKER})]
        rows.extend(_national_block(before))
        rows.append(_row({2: "Table 2: Regional breakdown"}))
        rows.extend(_national_block(after))
        rows.append(_row({2: _VI_MARKER}))
        rows.append(_row({}))
        rows.append(_row({}))
        rows.append(_row({col: name for name, col in _REGION_COLUMNS.items()}))
        rows.extend(_regional_block(_REGION_COLUMNS, _REGIONAL_FRACTIONS))

        workbook = build_workbook({"Tables": rows})
        with pytest.raises(ValueError, match="Missing expected party rows") as exc_info:
            bmg._parse_party_region_percentages(workbook)
        message = str(exc_info.value)
        assert "Conservative" not in message
        assert "Labour" not in message
        assert "Green" in message


# ── parse_poll ────────────────────────────────────────────────────────────────


class TestParsePoll:
    """Tests for parse_poll — combines fieldwork, sample and VI parsing."""

    def test_happy_path(self) -> None:
        workbook = _full_workbook()
        parsed = bmg.parse_poll(workbook, source_url=_XLSX_URL)

        assert parsed.sample_size == 1500
        assert parsed.fieldwork_start == date(2026, 1, 3)
        assert parsed.fieldwork_end == date(2026, 1, 4)
        assert len(parsed.party_region_percentages) == 8
        assert parsed.party_region_percentages["Labour"][bmg.NATIONAL_KEY] == 32.0

    def test_year_hint_used_when_url_and_text_have_no_year(self) -> None:
        workbook = _full_workbook(fieldwork_text="Fieldwork dates: 3-4 January")
        parsed = bmg.parse_poll(
            workbook, source_url="https://bmgresearch.com/tables.xlsx", year_hint=2026
        )
        assert parsed.fieldwork_start == date(2026, 1, 3)

    def test_no_year_and_no_hint_raises(self) -> None:
        workbook = _full_workbook(fieldwork_text="Fieldwork dates: 3-4 January")
        with pytest.raises(ValueError, match="Could not infer year"):
            bmg.parse_poll(workbook, source_url="https://bmgresearch.com/tables.xlsx")

    def test_propagates_missing_party_error(self) -> None:
        national = {
            name: pct for name, pct in _NATIONAL_FRACTIONS.items() if name != "Green"
        }
        workbook = _full_workbook(national=national)
        with pytest.raises(ValueError, match="Missing expected party rows"):
            bmg.parse_poll(workbook, source_url=_XLSX_URL)


# ── build_import_plan ──────────────────────────────────────────────────────────


class TestBuildImportPlan:
    """Tests for build_import_plan — dry-run plan from a fetched workbook."""

    def test_map_not_found_raises(
        self, db: Database, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        # A map-not-found plan must never reach the network: prove it by
        # making any fetch attempt fail the test outright, rather than
        # relying on the map check simply happening to run first.
        def fail_extract_workbook(*_a: object, **_k: object) -> Workbook:
            raise AssertionError("must not fetch when the map lookup already failed")

        monkeypatch.setattr(bmg, "extract_workbook", fail_extract_workbook)

        with pytest.raises(ValueError, match="Map not found: 'No Such Map'"):
            bmg.build_import_plan(db, map_name="No Such Map")

    def test_missing_parties_raises(
        self, db: Database, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        poll_map = db.add_map("A Map", parliament="westminster")

        def fake_extract_workbook(*_a: object, **_k: object) -> Workbook:
            return _full_workbook()

        monkeypatch.setattr(bmg, "extract_workbook", fake_extract_workbook)

        with pytest.raises(ValueError, match="Missing parties in database"):
            bmg.build_import_plan(db, map_name=poll_map.name)

    def test_one_row_per_db_region_with_zero_default(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        world = westminster_world

        def fake_extract_workbook(*_a: object, **_k: object) -> Workbook:
            return _full_workbook()

        monkeypatch.setattr(bmg, "extract_workbook", fake_extract_workbook)

        plan = bmg.build_import_plan(
            db,
            xlsx_url=_XLSX_URL,
            map_name=world.map_name,
            pollster_identifier="bmg_research",
        )

        assert plan.map_id == world.map_id
        assert plan.map_name == world.map_name
        assert plan.pollster_exists is False
        assert plan.pollster_id is None
        assert plan.pollster_name == "BMG Research"
        assert plan.regions_mapping == ""
        assert plan.poll_exists is False
        assert plan.poll_id is None

        # 8 parties × (1 national row + 12 regional rows).
        assert len(plan.rows) == 8 * 13
        national_rows = [row for row in plan.rows if row.region_id is None]
        assert len(national_rows) == 8

        labour_london = next(
            row
            for row in plan.rows
            if row.party_name == "Labour"
            and row.region_id == world.region_ids["London"]
        )
        assert labour_london.percentage == 42.0
        assert labour_london.party_id == world.party_ids["Labour"]

        labour_east_midlands = next(
            row
            for row in plan.rows
            if row.party_name == "Labour"
            and row.region_id == world.region_ids["East Midlands"]
        )
        assert labour_east_midlands.percentage == 0.0

        labour_national = next(
            row
            for row in plan.rows
            if row.party_name == "Labour" and row.region_id is None
        )
        assert labour_national.percentage == 32.0
        assert labour_national.region_name == "National"
        assert labour_national.party_id == world.party_ids["Labour"]

        # A second, distinct party's row, so a mutant that maps every row to
        # Labour's id (rather than each row's own party) cannot pass.
        conservative_national = next(
            row
            for row in plan.rows
            if row.party_name == "Conservative" and row.region_id is None
        )
        assert conservative_national.party_id == world.party_ids["Conservative"]

    def test_pollster_and_poll_existing_flags(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        world = westminster_world
        # Named differently from the code's "BMG Research" fallback default,
        # so a mutant that always returns the hard-coded default instead of
        # the seeded pollster's own name cannot pass.
        pollster = db.add_pollster("BMG Research Ltd", "bmg_research", weight=1.0)
        poll = db.add_poll(
            pollster.id,
            world.map_id,
            date(2026, 1, 3),
            date(2026, 1, 4),
            sample_size=1500,
        )

        def fake_extract_workbook(*_a: object, **_k: object) -> Workbook:
            return _full_workbook()

        monkeypatch.setattr(bmg, "extract_workbook", fake_extract_workbook)

        plan = bmg.build_import_plan(
            db,
            xlsx_url=_XLSX_URL,
            map_name=world.map_name,
            pollster_identifier="bmg_research",
        )

        assert plan.pollster_exists is True
        assert plan.pollster_id == pollster.id
        assert plan.pollster_name == "BMG Research Ltd"
        assert plan.poll_exists is True
        assert plan.poll_id == poll.id


# ── _cli_preview ────────────────────────────────────────────────────────────


class TestCliPreview:
    """Tests for _cli_preview — dry-run summary printed to stdout."""

    def _plan(self, **overrides: Any) -> bmg.ImportPlan:
        rows = overrides.pop(
            "rows",
            [
                bmg.PlannedPollRow(
                    party_id=1,
                    party_name="Labour",
                    region_id=None,
                    region_name="National",
                    percentage=32.0,
                ),
                bmg.PlannedPollRow(
                    party_id=1,
                    party_name="Labour",
                    region_id=5,
                    region_name="London",
                    percentage=42.0,
                ),
            ],
        )
        defaults: dict[str, Any] = {
            "pollster_identifier": "bmg_research",
            "pollster_name": "BMG Research",
            "pollster_id": None,
            "pollster_exists": False,
            "regions_mapping": "",
            "map_id": 1,
            "map_name": "UK Constituencies post 2022",
            "source_url": _XLSX_URL,
            "parsed": bmg.ParsedPoll(
                sample_size=1500,
                fieldwork_start=date(2026, 1, 3),
                fieldwork_end=date(2026, 1, 4),
                party_region_percentages={},
            ),
            "poll_id": None,
            "poll_exists": False,
            "rows": rows,
        }
        defaults.update(overrides)
        return bmg.ImportPlan(**defaults)

    def test_new_pollster_and_poll(self, capsys: pytest.CaptureFixture[str]) -> None:
        bmg._cli_preview(self._plan())
        out = capsys.readouterr().out

        assert "Parsed poll: fieldwork=2026-01-03 to 2026-01-04, sample=1500" in out
        assert "[dry-run] would create pollster: bmg_research" in out
        # "in out" would also match the pollster line above (a prefix of it),
        # so check the exact printed line instead.
        assert "[dry-run] would create poll" in out.splitlines()
        assert (
            "[dry-run] would insert row: party=Labour, region=National, "
            "region_id=None, pct=32.00" in out
        )
        assert (
            "[dry-run] would insert row: party=Labour, region=London, "
            "region_id=5, pct=42.00" in out
        )

    def test_existing_pollster_and_poll(
        self, capsys: pytest.CaptureFixture[str]
    ) -> None:
        bmg._cli_preview(self._plan(pollster_exists=True, poll_exists=True, poll_id=7))
        out = capsys.readouterr().out

        assert "pollster exists: bmg_research" in out
        assert "poll exists: 7" in out
        assert "[dry-run] would create pollster" not in out
        assert "[dry-run] would create poll" not in out


# ── main ────────────────────────────────────────────────────────────────────


class TestMain:
    """Tests for main — the CLI entry point, run against the ``db`` fixture."""

    def _run(
        self,
        db: Database,
        monkeypatch: pytest.MonkeyPatch,
        workbook: Workbook,
        *argv: str,
        fetched_urls: list[str] | None = None,
    ) -> None:
        """Run main() with Database and extract_workbook faked, via sys.argv.

        If ``fetched_urls`` is given, each URL ``extract_workbook`` is called
        with is appended to it, so a test can prove ``--xlsx-url`` actually
        reaches the fetch rather than the CLI default silently being used.
        """

        def fake_database(*_a: object, **_k: object) -> Database:
            return db

        def fake_extract_workbook(xlsx_url: str) -> Workbook:
            if fetched_urls is not None:
                fetched_urls.append(xlsx_url)
            return workbook

        monkeypatch.setattr(bmg, "Database", fake_database)
        monkeypatch.setattr(bmg, "extract_workbook", fake_extract_workbook)
        monkeypatch.setattr(sys, "argv", ["bmg_research_import.py", *argv])
        bmg.main()

    def test_dry_run_previews_without_writing(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        world = westminster_world
        self._run(
            db, monkeypatch, _full_workbook(), "--map-name", world.map_name, "--dry-run"
        )

        out = capsys.readouterr().out
        assert f"Fetching XLSX: {bmg.DEFAULT_XLSX_URL}" in out
        assert "[dry-run] would create pollster: bmg_research" in out
        assert len(db.get_all_pollsters()) == 0

    def test_commit_creates_pollster_poll_and_rows(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        world = westminster_world
        fetched_urls: list[str] = []
        self._run(
            db,
            monkeypatch,
            _full_workbook(),
            "--map-name",
            world.map_name,
            "--xlsx-url",
            "https://bmg.example.test/2026/x.xlsx",
            fetched_urls=fetched_urls,
        )

        out = capsys.readouterr().out
        assert "created pollster: bmg_research" in out
        assert "inserted poll rows: 104" in out

        pollsters = db.get_all_pollsters()
        assert len(pollsters) == 1
        assert pollsters[0].identifier == "bmg_research"
        polls = db.get_polls_by_pollster(pollsters[0].id)
        assert len(polls) == 1
        assert len(db.get_rows_for_poll(polls[0].id)) == 104
        assert fetched_urls == ["https://bmg.example.test/2026/x.xlsx"]
        assert polls[0].source_url == "https://bmg.example.test/2026/x.xlsx"

    def test_rerun_skips_then_replace_rows_overwrites(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        world = westminster_world
        argv = ("--map-name", world.map_name)

        self._run(db, monkeypatch, _full_workbook(), *argv)
        pollster = db.get_pollster_by_identifier("bmg_research")
        assert pollster is not None
        poll_id = db.get_polls_by_pollster(pollster.id)[0].id
        assert len(db.get_rows_for_poll(poll_id)) == 104

        self._run(db, monkeypatch, _full_workbook(), *argv)
        out = capsys.readouterr().out
        assert "poll exists" in out
        assert f"poll {poll_id} already has rows; use --replace-rows" in out
        assert len(db.get_rows_for_poll(poll_id)) == 104

        self._run(db, monkeypatch, _full_workbook(), *argv, "--replace-rows")
        out = capsys.readouterr().out
        assert "deleted existing rows: 104" in out
        assert "inserted poll rows: 104" in out
        assert len(db.get_rows_for_poll(poll_id)) == 104

    def test_year_hint_cli_flag(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        world = westminster_world
        workbook = _full_workbook(fieldwork_text="Fieldwork dates: 3-4 January")

        self._run(
            db,
            monkeypatch,
            workbook,
            "--map-name",
            world.map_name,
            "--xlsx-url",
            "https://bmgresearch.com/tables.xlsx",
            "--year-hint",
            "2026",
            "--dry-run",
        )

        out = capsys.readouterr().out
        assert "Parsed poll: fieldwork=2026-01-03 to 2026-01-04" in out

    def test_cli_defaults_pin(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        world = westminster_world
        assert bmg.DEFAULT_MAP_NAME == world.map_name
        assert bmg.DEFAULT_POLLSTER_IDENTIFIER == "bmg_research"

        self._run(db, monkeypatch, _full_workbook(), "--dry-run")

        out = capsys.readouterr().out
        assert "[dry-run] would create pollster: bmg_research" in out


class _PageResponse(FakeUrlResponse):
    def __init__(self, final_url: str, html: str) -> None:
        super().__init__(html.encode())
        self.final_url = final_url

    def geturl(self) -> str:
        return self.final_url

    def read(self, size: int = -1) -> bytes:
        return super().read()[:size] if size >= 0 else super().read()


_ARTICLE_URL = "https://inews.co.uk/news/politics/poll-478467"
_CANONICAL_ARTICLE_URL = "https://inews.co.uk/news/politics/poll-4784677"
_NEWS_URL = "https://bmgresearch.com/news/"
_RELEASE_URL = f"{_NEWS_URL}september-poll/"
_PUBLISHED_XLSX = (
    "https://bmgresearch.com/wp-content/uploads/2026/09/september-tables.xlsx"
)


def _links(*urls: str) -> str:
    return "".join(f'<a href="{url}">Source</a>' for url in urls)


def _mock_pages(
    monkeypatch: pytest.MonkeyPatch,
    pages: Mapping[str, tuple[str, str]],
) -> list[str]:
    requested: list[str] = []

    def fake_urlopen(req: Request, timeout: float = 45) -> _PageResponse:
        assert timeout == 45
        assert req.get_header("User-agent") == "Mozilla/5.0"
        assert req.get_header("Referer") == "https://bmgresearch.com/"
        requested.append(req.full_url)
        final_url, html = pages[req.full_url]
        return _PageResponse(final_url, html)

    monkeypatch.setattr(bmg, "urlopen", fake_urlopen)
    return requested


class TestResolveSourceUrl:
    @pytest.mark.parametrize(
        "url",
        [
            _XLSX_URL,
            "https://www.bmgresearch.co.uk/tables.xlsx",
            "https://example.com/legacy-download",
        ],
    )
    def test_direct_urls_do_not_fetch_html(
        self, monkeypatch: pytest.MonkeyPatch, url: str
    ) -> None:
        requested = _mock_pages(monkeypatch, {})
        assert bmg.resolve_source_url(url) == url
        assert requested == []

    def test_release_resolves_unique_relative_xlsx_link(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        _mock_pages(
            monkeypatch,
            {
                _RELEASE_URL: (
                    _RELEASE_URL,
                    _links(
                        "/wp-content/uploads/2026/09/september-tables.xlsx",
                        f"{_PUBLISHED_XLSX}#download",
                    ),
                ),
            },
        )
        assert bmg.resolve_source_url(_RELEASE_URL) == _PUBLISHED_XLSX

    def test_co_uk_release_remains_supported(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        source = "https://www.bmgresearch.co.uk/news/older-poll/"
        _mock_pages(monkeypatch, {source: (_RELEASE_URL, _links(_PUBLISHED_XLSX))})
        assert bmg.resolve_source_url(source) == _PUBLISHED_XLSX

    def test_original_article_redirect_requires_exact_canonical_backlink(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        other_release = f"{_NEWS_URL}different-poll/"
        requested = _mock_pages(
            monkeypatch,
            {
                _ARTICLE_URL: (_CANONICAL_ARTICLE_URL, "<html>Article</html>"),
                _NEWS_URL: (_NEWS_URL, _links(_RELEASE_URL, other_release)),
                _RELEASE_URL: (
                    _RELEASE_URL,
                    _links(f"{_CANONICAL_ARTICLE_URL}#results", _PUBLISHED_XLSX),
                ),
                other_release: (
                    other_release,
                    _links(_ARTICLE_URL, "/wp-content/uploads/2026/09/other.xlsx"),
                ),
            },
        )
        assert bmg.resolve_source_url(_ARTICLE_URL) == _PUBLISHED_XLSX
        assert requested == [_ARTICLE_URL, _NEWS_URL, _RELEASE_URL, other_release]

    def test_pagination_and_duplicate_links_are_bounded_and_visited_once(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        second_page = f"{_NEWS_URL}?results-page=2"
        requested = _mock_pages(
            monkeypatch,
            {
                _ARTICLE_URL: (_CANONICAL_ARTICLE_URL, ""),
                _NEWS_URL: (
                    _NEWS_URL,
                    _links("?results-page=2", "?results-page=2#next", _NEWS_URL),
                ),
                second_page: (
                    second_page,
                    _links(_NEWS_URL, _RELEASE_URL, _RELEASE_URL),
                ),
                _RELEASE_URL: (
                    _RELEASE_URL,
                    _links(_CANONICAL_ARTICLE_URL, _PUBLISHED_XLSX),
                ),
            },
        )
        assert bmg.resolve_source_url(_ARTICLE_URL) == _PUBLISHED_XLSX
        assert requested == [_ARTICLE_URL, _NEWS_URL, second_page, _RELEASE_URL]

    def test_news_traversal_ignores_foreign_hosts_and_non_news_paths(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        requested = _mock_pages(
            monkeypatch,
            {
                _ARTICLE_URL: (_CANONICAL_ARTICLE_URL, ""),
                _NEWS_URL: (
                    _NEWS_URL,
                    _links(
                        "https://bmgresearch.com.evil.example/news/poll/",
                        "https://bmgresearch.com@evil.example/news/poll/",
                        "https://evil.example/news/poll/",
                        "/about/",
                        "?news-category=polling",
                        _RELEASE_URL,
                    ),
                ),
                _RELEASE_URL: (
                    _RELEASE_URL,
                    _links(_CANONICAL_ARTICLE_URL, _PUBLISHED_XLSX, "/news/related/"),
                ),
            },
        )
        assert bmg.resolve_source_url(_ARTICLE_URL) == _PUBLISHED_XLSX
        assert requested == [_ARTICLE_URL, _NEWS_URL, _RELEASE_URL]

    @pytest.mark.parametrize(
        "links",
        [
            [],
            ["https://evil.example/wp-content/uploads/tables.xlsx"],
            ["https://bmgresearch.com.evil.example/wp-content/uploads/tables.xlsx"],
            ["https://bmgresearch.com/other/tables.xlsx"],
            [_PUBLISHED_XLSX, _PUBLISHED_XLSX.replace("september", "other")],
        ],
    )
    def test_release_requires_one_legitimate_published_workbook(
        self, monkeypatch: pytest.MonkeyPatch, links: list[str]
    ) -> None:
        _mock_pages(monkeypatch, {_RELEASE_URL: (_RELEASE_URL, _links(*links))})
        with pytest.raises(ValueError, match="Expected one published BMG XLSX"):
            bmg.resolve_source_url(_RELEASE_URL)

    def test_ambiguous_backlink_matches_fail(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        other_release = f"{_NEWS_URL}second-poll/"
        _mock_pages(
            monkeypatch,
            {
                _ARTICLE_URL: (_CANONICAL_ARTICLE_URL, ""),
                _NEWS_URL: (_NEWS_URL, _links(_RELEASE_URL, other_release)),
                _RELEASE_URL: (
                    _RELEASE_URL,
                    _links(_CANONICAL_ARTICLE_URL, _PUBLISHED_XLSX),
                ),
                other_release: (
                    other_release,
                    _links(
                        _CANONICAL_ARTICLE_URL,
                        _PUBLISHED_XLSX.replace("september", "different"),
                    ),
                ),
            },
        )
        with pytest.raises(ValueError, match="Multiple BMG workbooks"):
            bmg.resolve_source_url(_ARTICLE_URL)

    def test_repeated_matching_release_for_same_workbook_is_unambiguous(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        other_release = f"{_NEWS_URL}second-poll/"
        _mock_pages(
            monkeypatch,
            {
                _ARTICLE_URL: (_CANONICAL_ARTICLE_URL, ""),
                _NEWS_URL: (_NEWS_URL, _links(_RELEASE_URL, other_release)),
                _RELEASE_URL: (
                    _RELEASE_URL,
                    _links(_CANONICAL_ARTICLE_URL, _PUBLISHED_XLSX),
                ),
                other_release: (
                    other_release,
                    _links(_CANONICAL_ARTICLE_URL, _PUBLISHED_XLSX),
                ),
            },
        )
        assert bmg.resolve_source_url(_ARTICLE_URL) == _PUBLISHED_XLSX

    def test_no_canonical_backlink_fails_without_guessing(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        _mock_pages(
            monkeypatch,
            {
                _ARTICLE_URL: (_CANONICAL_ARTICLE_URL, ""),
                _NEWS_URL: (_NEWS_URL, _links(_RELEASE_URL)),
                _RELEASE_URL: (
                    _RELEASE_URL,
                    _links(_ARTICLE_URL, _PUBLISHED_XLSX),
                ),
            },
        )
        with pytest.raises(ValueError, match="No published BMG release"):
            bmg.resolve_source_url(_ARTICLE_URL)

    @pytest.mark.parametrize("kind", ["listing", "release"])
    def test_search_limits_fail_clearly(
        self, monkeypatch: pytest.MonkeyPatch, kind: str
    ) -> None:
        second_page = f"{_NEWS_URL}?results-page=2"
        next_release = f"{_NEWS_URL}second-poll/"
        pages = {
            _ARTICLE_URL: (_CANONICAL_ARTICLE_URL, ""),
            _NEWS_URL: (
                _NEWS_URL,
                _links(second_page)
                if kind == "listing"
                else _links(_RELEASE_URL, next_release),
            ),
            _RELEASE_URL: (_RELEASE_URL, ""),
        }
        limit = "_MAX_RELEASE_PAGES" if kind == "release" else "_MAX_NEWS_PAGES"
        monkeypatch.setattr(bmg, limit, 1)
        requested = _mock_pages(monkeypatch, pages)
        with pytest.raises(ValueError, match="search limit reached"):
            bmg.resolve_source_url(_ARTICLE_URL)
        assert second_page not in requested
        assert next_release not in requested

    @pytest.mark.parametrize("source", [_ARTICLE_URL, _RELEASE_URL])
    def test_foreign_redirect_is_rejected(
        self, monkeypatch: pytest.MonkeyPatch, source: str
    ) -> None:
        _mock_pages(monkeypatch, {source: ("https://evil.example/news/poll/", "")})
        with pytest.raises(ValueError, match="Unexpected source-page redirect"):
            bmg.resolve_source_url(source)

    def test_non_news_listing_redirect_is_rejected(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        _mock_pages(
            monkeypatch,
            {
                _ARTICLE_URL: (_CANONICAL_ARTICLE_URL, ""),
                _NEWS_URL: ("https://bmgresearch.com/about/", ""),
            },
        )
        with pytest.raises(ValueError, match="Unexpected BMG news listing redirect"):
            bmg.resolve_source_url(_ARTICLE_URL)

    def test_oversize_html_fails_clearly(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        _mock_pages(monkeypatch, {_RELEASE_URL: (_RELEASE_URL, "x" * 2_000_001)})
        with pytest.raises(ValueError, match="Source page is too large"):
            bmg.resolve_source_url(_RELEASE_URL)


def _staged_vi_workbook(
    final: Mapping[str, float] = _NATIONAL_FRACTIONS,
    *,
    base: str = "Base: Not sure and prefer not to say removed",
) -> Workbook:
    rows = [
        _row({2: _VI_MARKER}),
        *_national_block({party: 0.01 for party in final}),
    ]
    rows.extend([_row({2: "Table 2"}), _row({2: _VI_MARKER}), _row({2: base})])
    rows.extend(_national_block(final))
    rows.append(_row({2: "Table 3: unrelated party question"}))
    rows.extend(_national_block({party: 0.99 for party in _REQUIRED_PARTIES}))
    return build_workbook(
        {
            "Methodology": _methodology_sheet(
                sample_text="Sample: 1,515 GB adults aged 18+"
            ),
            "Tables": _wrap_with_regional_trailer(rows),
        }
    )


class TestFinalHeadlineAndSample:
    @pytest.mark.parametrize(
        ("sample_text", "expected"),
        [
            ("Sample: 1,515 GB adults aged 18+", 1515),
            ("Sample: 1579 GB adults aged 18 to 75", 1579),
            ("Sample: 1,559 adults in 11 regions", 1559),
        ],
    )
    def test_sample_count_excludes_demographic_digits(
        self, sample_text: str, expected: int
    ) -> None:
        workbook = build_workbook(
            {"Methodology": _methodology_sheet(sample_text=sample_text)}
        )
        assert bmg._parse_fieldwork_and_sample(workbook, 2026)[2] == expected

    @pytest.mark.parametrize(
        "sample_text",
        [
            "Sample: GB adults 18+",
            "Sample: 0 adults aged 18+",
            "Sample: 1,xxx GB adults aged 18+",
            "Sample: 1,51 GB adults aged 18+",
            "Sample: 1,515, GB adults aged 18+",
        ],
    )
    def test_age_digits_do_not_supply_missing_sample(self, sample_text: str) -> None:
        workbook = build_workbook(
            {"Methodology": _methodology_sheet(sample_text=sample_text)}
        )
        with pytest.raises(ValueError, match="Could not parse sample size"):
            bmg._parse_fieldwork_and_sample(workbook, 2026)

    @pytest.mark.parametrize(
        "base",
        [
            "Base: Not sure and prefer not to say removed",
            "BASE: Not sure   and prefer not to say removed\u00a0",
        ],
    )
    def test_explicit_final_base_selects_headline_without_changing_regions(
        self, base: str
    ) -> None:
        parsed = bmg.parse_poll(_staged_vi_workbook(base=base), source_url=_XLSX_URL)
        legacy = bmg.parse_poll(_full_workbook(), source_url=_XLSX_URL)
        assert parsed.sample_size == 1515
        assert {
            party: values[bmg.NATIONAL_KEY]
            for party, values in parsed.party_region_percentages.items()
        } == {
            party: round(value * 100, 2)
            for party, value in _NATIONAL_FRACTIONS.items()
        }
        assert {
            party: {
                region: value
                for region, value in values.items()
                if region != bmg.NATIONAL_KEY
            }
            for party, values in parsed.party_region_percentages.items()
        } == {
            party: {
                region: value
                for region, value in values.items()
                if region != bmg.NATIONAL_KEY
            }
            for party, values in legacy.party_region_percentages.items()
        }

    def test_duplicate_explicit_final_tables_fail(self) -> None:
        workbook = _staged_vi_workbook()
        sheet = workbook["Tables"]
        sheet.append(_row({2: _VI_MARKER}))
        sheet.append(_row({2: "Base: Not sure and prefer not to say removed"}))
        with pytest.raises(ValueError, match="Multiple final voting-intention"):
            bmg._parse_party_region_percentages(workbook)

    def test_next_table_cannot_supply_a_missing_headline_party(self) -> None:
        final = {
            party: value
            for party, value in _NATIONAL_FRACTIONS.items()
            if party != "Green"
        }
        with pytest.raises(ValueError, match="Missing expected party rows.*Green"):
            bmg._parse_party_region_percentages(_staged_vi_workbook(final))

    def test_next_vi_marker_does_not_overwrite_legacy_nationals(self) -> None:
        workbook = _full_workbook()
        starts = bmg._find_vi_table_starts(workbook["Tables"])
        workbook["Tables"].cell(starts[-1] + 4, 3, 0.99)
        parsed = bmg._parse_party_region_percentages(workbook)
        assert parsed["Labour"][bmg.NATIONAL_KEY] == 32.0

    def test_early_next_table_cannot_supply_a_missing_final_headline(self) -> None:
        rows = [
            _row({2: _VI_MARKER}),
            _row({2: "Base: Not sure and prefer not to say removed"}),
            _row({2: "Table 2: unrelated"}),
            *_national_block(_NATIONAL_FRACTIONS),
        ]
        workbook = build_workbook({"Tables": _wrap_with_regional_trailer(rows)})
        with pytest.raises(ValueError, match="Missing expected party rows"):
            bmg._parse_party_region_percentages(workbook)

    def test_resolved_workbook_url_controls_year_and_provenance(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        resolved_url = _PUBLISHED_XLSX.replace("2026", "2025")
        requested: list[str] = []

        def fake_resolve(url: str) -> str:
            assert url == _ARTICLE_URL
            return resolved_url

        def fake_extract(url: str) -> Workbook:
            requested.append(url)
            return _full_workbook(fieldwork_text="Fieldwork dates: 3-4 January")

        monkeypatch.setattr(bmg, "resolve_source_url", fake_resolve)
        monkeypatch.setattr(bmg, "extract_workbook", fake_extract)
        plan = bmg.build_import_plan(db, xlsx_url=_ARTICLE_URL, year_hint=2026)
        assert plan.source_url == resolved_url
        assert plan.parsed.fieldwork_start == date(2025, 1, 3)
        assert plan.parsed.fieldwork_end == date(2025, 1, 4)
        assert requested == [resolved_url]
