"""Tests for the Focaldata XLSX / Google Sheets poll importer.

Covers parsing helpers, ``build_import_plan``, ``_cli_preview`` and ``main``.
``commit_import_plan`` and ``_find_existing_poll`` are already covered for every
Westminster importer, including this one, by
``test_westminster_importers_commit.py`` and are not retested here.
"""

from __future__ import annotations

import sys
from collections.abc import Mapping
from datetime import date
from types import MappingProxyType
from typing import Any
from urllib.request import Request

import pytest
from openpyxl import Workbook

from db import Database
from polls.importers.westminster import focaldata_import as fd
from tests.uk_fixtures import (
    FakeUrlResponse,
    WestminsterWorld,
    build_workbook,
    workbook_bytes,
)

# ── Synthetic workbook builders ────────────────────────────────────────────────
#
# The "Tables" sheet layout the parser expects: a voting-intention header row
# (found by _find_vi_section_start), a region-header row listing macro-region
# names in columns 11+ (found by scanning columns 11-24), then one data row per
# party: the party label in column 1, its national percentage in column 2, and
# its per-macro-region percentages in the macro-region columns.

_COL_WIDTH = 20
_HEADER_PHRASE = "Combined voting intention (excluding don't knows)"

_MACRO_COLUMNS: Mapping[str, int] = MappingProxyType(
    {
        "North of England": 11,
        "Midlands": 12,
        "South of England": 13,
        "Greater London": 14,
        "Wales": 15,
        "Scotland": 16,
    }
)

# Fractions in [0.0, 1.0]; _to_percentage multiplies by 100. Sums to 1.0.
_NATIONAL_FRACTIONS: Mapping[str, float] = MappingProxyType(
    {
        "Conservative": 0.24,
        "Labour": 0.32,
        "Liberal Democrats": 0.11,
        "Reform UK": 0.21,
        "Green": 0.06,
        "Scottish National Party": 0.03,
        "Plaid Cymru": 0.01,
        "Other": 0.02,
    }
)

_REQUIRED_PARTIES: tuple[str, ...] = tuple(_NATIONAL_FRACTIONS.keys())

# Only Labour and Conservative get regional macro figures, so every other
# party's macro columns come out at the 0.0 default.
_LABOUR_MACROS: Mapping[str, float] = MappingProxyType(
    {
        "North of England": 0.38,
        "Midlands": 0.34,
        "South of England": 0.25,
        "Greater London": 0.45,
        "Wales": 0.40,
        "Scotland": 0.05,
    }
)
_CONSERVATIVE_MACROS: Mapping[str, float] = MappingProxyType(
    {
        "North of England": 0.20,
        "Midlands": 0.26,
        "South of England": 0.30,
        "Greater London": 0.18,
        "Wales": 0.22,
        "Scotland": 0.15,
    }
)
_PARTY_MACROS: Mapping[str, Mapping[str, float]] = MappingProxyType(
    {"Labour": _LABOUR_MACROS, "Conservative": _CONSERVATIVE_MACROS}
)

_FIELDWORK_TEXT = "16-19 Jan 2026"
_SAMPLE_SIZE = 2032
_XLSX_URL = "https://focaldata.example.test/2026/tables.xlsx"

# Appended to the end of every hand-built "Tables" sheet in this file.
# focaldata_import.py:624 scans a range with an EXCLUSIVE upper bound
# (`range(start_row, min(sheet.max_row, start_row + 180))`), so whatever row
# sits at the sheet's actual `max_row` is silently never read. Without a
# trailer, the last party appended by `_tables_sheet` (always "Other", the
# last key in `_NATIONAL_FRACTIONS`) would land exactly on that excluded row
# and its value would silently come from the optional-party defaulting
# fallback instead of the real parsed figure — see
# test_party_on_sheets_actual_last_row_is_silently_dropped_pins_current_behaviour,
# which deliberately omits this trailer to demonstrate the bug.
_TRAILER_ROW_TEXT = "n=2,032 GB adults"


def _row(values: Mapping[int, object]) -> list[object]:
    """One 1-indexed worksheet row (as a 0-indexed list), padded to _COL_WIDTH."""
    cells: list[object] = [None] * _COL_WIDTH
    for col, value in values.items():
        cells[col - 1] = value
    return cells


def _party_row(
    party: str, fraction: float, macros: Mapping[str, float] | None = None
) -> list[object]:
    """One party's data row: label in col 1, national % in col 2, and its
    macro-region %s in their own columns.
    """
    values: dict[int, object] = {1: party, 2: fraction}
    for macro_name, pct in (macros or {}).items():
        values[_MACRO_COLUMNS[macro_name]] = pct
    return _row(values)


def _header_rows(
    *, macro_columns: Mapping[str, int] | None = _MACRO_COLUMNS
) -> list[list[object]]:
    """The VI question header row, plus the region-header row.

    ``macro_columns=None`` omits the region-header row, to test the
    "region header not found" path.
    """
    rows: list[list[object]] = [_row({1: _HEADER_PHRASE})]
    if macro_columns is not None:
        rows.append(_row({col: name for name, col in macro_columns.items()}))
    return rows


def _tables_sheet(
    *,
    national: Mapping[str, float] = _NATIONAL_FRACTIONS,
    macros: Mapping[str, Mapping[str, float]] = _PARTY_MACROS,
    macro_columns: Mapping[str, int] | None = _MACRO_COLUMNS,
    extra_rows: list[list[object]] | None = None,
) -> list[list[object]]:
    """A "Tables" sheet: header rows, one row per party, a harmless trailer.

    ``extra_rows``, if given, are inserted between the party rows and the
    mandatory trailer row — e.g. a "Column N=..." break-condition row
    followed by rows that should be ignored because of it.
    """
    rows = _header_rows(macro_columns=macro_columns)
    for party, fraction in national.items():
        rows.append(_party_row(party, fraction, macros.get(party)))
    if extra_rows:
        rows.extend(extra_rows)
    rows.append(_row({1: _TRAILER_ROW_TEXT}))
    return rows


def _info_sheet(
    *,
    fieldwork_label: str | None = "Dates conducted",
    fieldwork_value: object = _FIELDWORK_TEXT,
) -> list[list[object]]:
    """An "Info" sheet with a "dates conducted" row and a "sample size" row."""
    rows: list[list[object]] = [_row({1: "Focaldata Omnibus"})]
    if fieldwork_label is not None:
        rows.append(_row({2: fieldwork_label, 3: fieldwork_value}))
    rows.append(_row({2: "Sample size", 3: _SAMPLE_SIZE}))
    return rows


def _full_workbook(
    *,
    fieldwork_text: str | None = _FIELDWORK_TEXT,
    national: Mapping[str, float] = _NATIONAL_FRACTIONS,
    macros: Mapping[str, Mapping[str, float]] = _PARTY_MACROS,
) -> Workbook:
    """A full workbook: an "Info" sheet plus a "Tables" sheet."""
    return build_workbook(
        {
            "Info": _info_sheet(fieldwork_value=fieldwork_text),
            "Tables": _tables_sheet(national=national, macros=macros),
        }
    )


# ── _month_number ────────────────────────────────────────────────────────────


class TestMonthNumber:
    """Tests for _month_number — month name → integer conversion."""

    def test_full_names(self) -> None:
        assert fd._month_number("January") == 1
        assert fd._month_number("February") == 2
        assert fd._month_number("March") == 3
        assert fd._month_number("April") == 4
        assert fd._month_number("May") == 5
        assert fd._month_number("June") == 6
        assert fd._month_number("July") == 7
        assert fd._month_number("August") == 8
        assert fd._month_number("September") == 9
        assert fd._month_number("October") == 10
        assert fd._month_number("November") == 11
        assert fd._month_number("December") == 12

    def test_abbreviated_names(self) -> None:
        assert fd._month_number("Jan") == 1
        assert fd._month_number("Feb") == 2
        assert fd._month_number("Sep") == 9
        assert fd._month_number("Sept") == 9
        assert fd._month_number("Dec") == 12

    def test_case_insensitive(self) -> None:
        assert fd._month_number("JANUARY") == 1
        assert fd._month_number("january") == 1
        assert fd._month_number("jAnUaRy") == 1

    def test_trailing_period_stripped(self) -> None:
        assert fd._month_number("Jan.") == 1

    def test_unknown_returns_none(self) -> None:
        assert fd._month_number("Octember") is None
        assert fd._month_number("") is None
        assert fd._month_number("13") is None


# ── _infer_year ───────────────────────────────────────────────────────────────


class TestInferYear:
    """Tests for _infer_year — four-digit year extraction from URLs."""

    def test_year_as_path_segment(self) -> None:
        url = "https://focaldata.example.test/2026/tables.xlsx"
        assert fd._infer_year(url) == 2026

    def test_year_as_bare_occurrence(self) -> None:
        url = "https://focaldata.example.test/tables_2025_v2.xlsx"
        assert fd._infer_year(url) == 2025

    def test_path_segment_preferred_over_bare(self) -> None:
        url = "https://focaldata.example.test/2024archive/2026/tables.xlsx"
        assert fd._infer_year(url) == 2026

    def test_fallback_when_no_year_in_url(self) -> None:
        url = "https://focaldata.example.test/tables.xlsx"
        assert fd._infer_year(url, fallback=2025) == 2025

    def test_no_year_no_fallback_raises(self) -> None:
        with pytest.raises(ValueError, match="Could not infer year"):
            fd._infer_year("https://focaldata.example.test/tables.xlsx")


# ── _parse_fieldwork ──────────────────────────────────────────────────────────


class TestParseFieldwork:
    """Tests for _parse_fieldwork — date-range string parsing."""

    def test_same_month_with_year(self) -> None:
        start, end = fd._parse_fieldwork("16-19 Jan 2026")
        assert start == date(2026, 1, 16)
        assert end == date(2026, 1, 19)

    def test_same_month_no_year_uses_default(self) -> None:
        start, end = fd._parse_fieldwork("16-19 Jan", default_year=2026)
        assert start == date(2026, 1, 16)
        assert end == date(2026, 1, 19)

    def test_cross_month_with_year(self) -> None:
        start, end = fd._parse_fieldwork("31 Jan - 3 Feb 2026")
        assert start == date(2026, 1, 31)
        assert end == date(2026, 2, 3)

    def test_cross_month_no_year_uses_default(self) -> None:
        start, end = fd._parse_fieldwork("31 Jan - 3 Feb", default_year=2026)
        assert start == date(2026, 1, 31)
        assert end == date(2026, 2, 3)

    def test_cross_year_with_year_decrements_start_year(self) -> None:
        start, end = fd._parse_fieldwork("31 Dec - 2 Jan 2026")
        assert start == date(2025, 12, 31)
        assert end == date(2026, 1, 2)

    def test_cross_year_no_year_decrements_start_year(self) -> None:
        start, end = fd._parse_fieldwork("31 Dec - 2 Jan", default_year=2026)
        assert start == date(2025, 12, 31)
        assert end == date(2026, 1, 2)

    def test_ordinal_suffixes_stripped(self) -> None:
        start, end = fd._parse_fieldwork("16th-19th Jan 2026")
        assert start == date(2026, 1, 16)
        assert end == date(2026, 1, 19)

    @pytest.mark.parametrize("dash", ["–", "—"], ids=["en_dash", "em_dash"])
    def test_dash_variants_normalised(self, dash: str) -> None:
        start, end = fd._parse_fieldwork(f"31 Jan {dash} 3 Feb 2026")
        assert start == date(2026, 1, 31)
        assert end == date(2026, 2, 3)

    def test_no_pattern_matches_raises(self) -> None:
        with pytest.raises(ValueError, match="Could not parse fieldwork string"):
            fd._parse_fieldwork("not a date")

    def test_no_year_pattern_without_default_year_falls_through(self) -> None:
        # "16-19 Jan" without default_year matches no pattern: the no-year
        # pattern requires default_year, and nothing else fits.
        with pytest.raises(ValueError, match="Could not parse fieldwork string"):
            fd._parse_fieldwork("16-19 Jan")

    def test_same_month_unrecognised_month_raises(self) -> None:
        with pytest.raises(ValueError, match="Could not parse month in fieldwork"):
            fd._parse_fieldwork("16-19 Blorpuary 2026")

    def test_cross_month_unrecognised_month_raises(self) -> None:
        with pytest.raises(ValueError, match="Could not parse months in fieldwork"):
            fd._parse_fieldwork("31 Blorpuary - 3 Feb 2026")

    def test_same_month_no_year_unrecognised_month_raises(self) -> None:
        with pytest.raises(ValueError, match="Could not parse month in fieldwork"):
            fd._parse_fieldwork("16-19 Blorpuary", default_year=2026)

    def test_cross_month_no_year_unrecognised_month_raises(self) -> None:
        with pytest.raises(ValueError, match="Could not parse months in fieldwork"):
            fd._parse_fieldwork("31 Blorpuary - 3 Feb", default_year=2026)


# ── _google_sheet_export_url ──────────────────────────────────────────────────


class TestGoogleSheetExportUrl:
    """Tests for _google_sheet_export_url — Sheets URL → export URL, else None."""

    def test_sheets_url_with_gid(self) -> None:
        url = "https://docs.google.com/spreadsheets/d/ABC123/edit?gid=456"
        assert fd._google_sheet_export_url(url) == (
            "https://docs.google.com/spreadsheets/d/ABC123/export?format=xlsx&gid=456"
        )

    def test_sheets_url_without_gid(self) -> None:
        url = "https://docs.google.com/spreadsheets/d/ABC123/edit?usp=sharing"
        assert fd._google_sheet_export_url(url) == (
            "https://docs.google.com/spreadsheets/d/ABC123/export?format=xlsx"
        )

    def test_non_google_url_returns_none(self) -> None:
        assert fd._google_sheet_export_url("https://focaldata.example.test/x.xlsx") is None

    def test_google_domain_wrong_path_returns_none(self) -> None:
        url = "https://docs.google.com/document/d/ABC123/edit"
        assert fd._google_sheet_export_url(url) is None

    def test_spreadsheets_path_without_id_returns_none(self) -> None:
        # The domain/path substring check passes, but nothing follows the
        # "/spreadsheets/d/" prefix for the id regex to capture.
        url = "https://docs.google.com/spreadsheets/d/"
        assert fd._google_sheet_export_url(url) is None


# ── extract_workbook ────────────────────────────────────────────────────────


def _valid_payload() -> bytes:
    """XLSX bytes for a trivial one-sheet workbook (starts with the PK magic)."""
    return workbook_bytes(build_workbook({"Sheet1": [["ok"]]}))


class TestExtractWorkbook:
    """Tests for extract_workbook — fetch (and possibly convert) an XLSX URL."""

    def test_plain_url_success(self, monkeypatch: pytest.MonkeyPatch) -> None:
        payload = _valid_payload()

        def fake_urlopen(req: Request, timeout: float = 50) -> FakeUrlResponse:
            assert req.full_url == "https://focaldata.example.test/x.xlsx"
            return FakeUrlResponse(payload)

        monkeypatch.setattr(fd, "urlopen", fake_urlopen)
        workbook = fd.extract_workbook("https://focaldata.example.test/x.xlsx")
        assert workbook.sheetnames == ["Sheet1"]

    def test_google_sheets_url_tries_export_first_and_succeeds(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        payload = _valid_payload()
        requested: list[str] = []
        original = "https://docs.google.com/spreadsheets/d/ABC123/edit?gid=1"
        export = "https://docs.google.com/spreadsheets/d/ABC123/export?format=xlsx&gid=1"

        def fake_urlopen(req: Request, timeout: float = 50) -> FakeUrlResponse:
            requested.append(req.full_url)
            return FakeUrlResponse(payload)

        monkeypatch.setattr(fd, "urlopen", fake_urlopen)
        workbook = fd.extract_workbook(original)

        assert workbook.sheetnames == ["Sheet1"]
        # Only the export URL is requested: the loop breaks on first success.
        assert requested == [export]

    def test_google_sheets_export_fails_falls_back_to_original(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        payload = _valid_payload()
        requested: list[str] = []
        original = "https://docs.google.com/spreadsheets/d/ABC123/edit?gid=1"
        export = "https://docs.google.com/spreadsheets/d/ABC123/export?format=xlsx&gid=1"

        def fake_urlopen(req: Request, timeout: float = 50) -> FakeUrlResponse:
            requested.append(req.full_url)
            if req.full_url == export:
                return FakeUrlResponse(b"<html>not xlsx</html>")
            return FakeUrlResponse(payload)

        monkeypatch.setattr(fd, "urlopen", fake_urlopen)
        workbook = fd.extract_workbook(original)

        assert workbook.sheetnames == ["Sheet1"]
        assert requested == [export, original]

    def test_non_xlsx_payload_raises(self, monkeypatch: pytest.MonkeyPatch) -> None:
        def fake_urlopen(req: Request, timeout: float = 50) -> FakeUrlResponse:
            return FakeUrlResponse(b"<html>not an xlsx file</html>")

        monkeypatch.setattr(fd, "urlopen", fake_urlopen)
        with pytest.raises(ValueError, match="non-xlsx payload"):
            fd.extract_workbook("https://focaldata.example.test/x.xlsx")

    def test_urlopen_exception_is_captured_in_message(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        def fake_urlopen(req: Request, timeout: float = 50) -> FakeUrlResponse:
            raise TimeoutError("connection timed out")

        monkeypatch.setattr(fd, "urlopen", fake_urlopen)
        with pytest.raises(ValueError, match="connection timed out"):
            fd.extract_workbook("https://focaldata.example.test/x.xlsx")

    def test_all_candidates_failing_combines_errors(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        original = "https://docs.google.com/spreadsheets/d/ABC123/edit?gid=1"

        def fake_urlopen(req: Request, timeout: float = 50) -> FakeUrlResponse:
            raise ConnectionError("unreachable")

        monkeypatch.setattr(fd, "urlopen", fake_urlopen)
        with pytest.raises(ValueError, match="Could not fetch XLSX payload") as exc_info:
            fd.extract_workbook(original)
        message = str(exc_info.value)
        assert message.count("unreachable") == 2
        assert " | " in message


# ── _cell_text ────────────────────────────────────────────────────────────────


class TestCellText:
    """Tests for _cell_text — openpyxl cell value → stripped string."""

    def test_none_returns_empty_string(self) -> None:
        assert fd._cell_text(None) == ""

    def test_string_stripped(self) -> None:
        assert fd._cell_text("  hello  ") == "hello"

    def test_integer_converted(self) -> None:
        assert fd._cell_text(42) == "42"

    def test_float_converted(self) -> None:
        assert fd._cell_text(3.14) == "3.14"

    def test_empty_string(self) -> None:
        assert fd._cell_text("") == ""


# ── _to_percentage / _to_percentage_or_zero ────────────────────────────────────


class TestToPercentage:
    """Tests for _to_percentage — raw cell value → percentage float."""

    def test_integer_in_0_100_range(self) -> None:
        assert fd._to_percentage(42) == pytest.approx(42.0)

    def test_decimal_in_0_1_range_multiplied(self) -> None:
        assert fd._to_percentage(0.42) == pytest.approx(42.0)

    def test_rounds_to_nearest_integer(self) -> None:
        assert fd._to_percentage(34.6) == pytest.approx(35.0)
        assert fd._to_percentage(34.4) == pytest.approx(34.0)

    def test_zero(self) -> None:
        assert fd._to_percentage(0) == pytest.approx(0.0)

    def test_one_boundary_treated_as_100_percent(self) -> None:
        assert fd._to_percentage(1.0) == pytest.approx(100.0)

    def test_returns_a_float(self) -> None:
        assert type(fd._to_percentage(42)) is float

    def test_none_raises(self) -> None:
        with pytest.raises(ValueError, match="empty percentage cell"):
            fd._to_percentage(None)


class TestToPercentageOrZero:
    """Tests for _to_percentage_or_zero — blank-tolerant percentage conversion."""

    def test_none_returns_zero(self) -> None:
        assert fd._to_percentage_or_zero(None) == pytest.approx(0.0)

    def test_hyphen_returns_zero(self) -> None:
        assert fd._to_percentage_or_zero("-") == pytest.approx(0.0)

    def test_en_dash_returns_zero(self) -> None:
        assert fd._to_percentage_or_zero("–") == pytest.approx(0.0)

    def test_em_dash_returns_zero(self) -> None:
        assert fd._to_percentage_or_zero("—") == pytest.approx(0.0)

    def test_empty_string_returns_zero(self) -> None:
        assert fd._to_percentage_or_zero("") == pytest.approx(0.0)

    def test_numeric_delegates_to_to_percentage(self) -> None:
        assert fd._to_percentage_or_zero(0.35) == pytest.approx(35.0)

    def test_non_dash_string_delegates_to_to_percentage(self) -> None:
        assert fd._to_percentage_or_zero("42") == pytest.approx(42.0)

    def test_integer_in_0_100_range(self) -> None:
        assert fd._to_percentage_or_zero(28) == pytest.approx(28.0)


# ── _canonical_party ─────────────────────────────────────────────────────────


class TestCanonicalParty:
    """Tests for _canonical_party — raw label → canonical party name, or None."""

    @pytest.mark.parametrize(
        ("label", "expected"),
        [
            ("Conservative", "Conservative"),
            ("Labour", "Labour"),
            ("Liberal Democrats", "Liberal Democrats"),
            ("Liberal Democrat", "Liberal Democrats"),
            ("Reform UK", "Reform UK"),
            ("Green", "Green"),
            ("Green Party", "Green"),
            ("Scottish National Party", "Scottish National Party"),
            ("Scottish National Party (SNP)", "Scottish National Party"),
            ("SNP", "Scottish National Party"),
            ("Plaid Cymru", "Plaid Cymru"),
            ("Another party", "Other"),
            ("Other", "Other"),
            (
                "An independent candidate or other party (e.g. Workers Party "
                "/ SDP / Yorkshire Party)",
                "Other",
            ),
            (
                "An independent candidate or other party (e.g. Your Party / "
                "Workers Party / SDP)",
                "Other",
            ),
            ("Another party (e.g. Workers Party / SDP / Yorkshire Party)", "Other"),
        ],
        ids=lambda v: v if isinstance(v, str) and len(v) < 40 else None,
    )
    def test_known_labels_resolve(self, label: str, expected: str) -> None:
        # Every one of these labels also happens to resolve via a heuristic
        # branch below, so this doesn't prove PARTY_NAME_MAP's direct lookup
        # specifically fires (that lookup is a redundant fast path); it only
        # proves each label resolves to the right canonical party somehow.
        assert fd._canonical_party(label) == expected

    def test_heuristic_reform(self) -> None:
        assert fd._canonical_party("Reform Party UK") == "Reform UK"

    def test_heuristic_liberal_democrat(self) -> None:
        assert fd._canonical_party("Liberal Democrat Focus Team") == "Liberal Democrats"

    def test_heuristic_green(self) -> None:
        assert fd._canonical_party("Green Alliance") == "Green"

    def test_heuristic_conservative(self) -> None:
        assert fd._canonical_party("Conservative and Unionist Party") == "Conservative"

    def test_heuristic_labour_prefix(self) -> None:
        assert fd._canonical_party("Labour Co-operative") == "Labour"

    def test_heuristic_labour_not_at_start_returns_none(self) -> None:
        # "labour" must be a prefix, not merely present anywhere in the label.
        assert fd._canonical_party("New Labour") is None

    def test_heuristic_snp_substring(self) -> None:
        assert fd._canonical_party("Scottish National candidate") == (
            "Scottish National Party"
        )

    def test_heuristic_snp_exact_lowercase(self) -> None:
        # "SNP" (upper) is a direct PARTY_NAME_MAP hit; lowercase "snp" is not,
        # so it must fall through to the case-insensitive heuristic instead.
        assert fd._canonical_party("snp") == "Scottish National Party"

    def test_heuristic_plaid(self) -> None:
        assert fd._canonical_party("Plaid Cymru Wales") == "Plaid Cymru"

    def test_heuristic_independent(self) -> None:
        assert fd._canonical_party("An Independent") == "Other"

    def test_heuristic_another_party_substring(self) -> None:
        assert fd._canonical_party("Yet another party entirely") == "Other"

    def test_heuristic_other_exact_lowercase(self) -> None:
        # "Other" (title case) is a direct hit; "OTHER" is not, so it must
        # fall through to the case-insensitive heuristic.
        assert fd._canonical_party("OTHER") == "Other"

    def test_unmatched_returns_none(self) -> None:
        assert fd._canonical_party("Workers Party") is None

    def test_blank_label_returns_none(self) -> None:
        assert fd._canonical_party("") is None


# ── _parse_info_sheet ──────────────────────────────────────────────────────────


class TestParseInfoSheet:
    """Tests for _parse_info_sheet — fieldwork/sample extraction, with fallbacks."""

    def test_happy_path_via_b_c_columns(self) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    [None, None, None],
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    [None, "Sample size", 2032],
                ]
            }
        )
        start, end, sample = fd._parse_info_sheet(workbook, 2026)
        assert start == date(2026, 1, 16)
        assert end == date(2026, 1, 19)
        assert sample == 2032

    def test_falls_back_to_first_sheet_when_no_info_sheet(self) -> None:
        workbook = build_workbook(
            {
                "Cover": [
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    [None, "Sample size", 2032],
                ]
            }
        )
        start, end, sample = fd._parse_info_sheet(workbook, 2026)
        assert start == date(2026, 1, 16)
        assert sample == 2032

    def test_last_dates_conducted_match_wins(self) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    [None, "Dates conducted", "1-2 Jan 2026"],
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    [None, "Sample size", 2032],
                ]
            }
        )
        start, end, _sample = fd._parse_info_sheet(workbook, 2026)
        assert start == date(2026, 1, 16)
        assert end == date(2026, 1, 19)

    def test_last_sample_size_match_wins(self) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    [None, "Sample size", 999],
                    [None, "Sample size", 2032],
                ]
            }
        )
        _start, _end, sample = fd._parse_info_sheet(workbook, 2026)
        assert sample == 2032

    def test_sample_size_string_with_digits(self) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    [None, "Sample size", "2,032 respondents"],
                ]
            }
        )
        _start, _end, sample = fd._parse_info_sheet(workbook, 2026)
        assert sample == 2032

    def test_sample_size_float_rounded(self) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    [None, "Sample size", 2031.6],
                ]
            }
        )
        _start, _end, sample = fd._parse_info_sheet(workbook, 2026)
        assert sample == 2032

    def test_fieldwork_a_b_fallback_col_b_populated(self) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    ["Dates conducted", "16-19 Jan 2026"],
                    [None, "Sample size", 2032],
                ]
            }
        )
        start, end, _sample = fd._parse_info_sheet(workbook, 2026)
        assert start == date(2026, 1, 16)
        assert end == date(2026, 1, 19)

    def test_fieldwork_a_b_fallback_col_b_empty_uses_col_c(self) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    ["Dates conducted", None, "16-19 Jan 2026"],
                    [None, "Sample size", 2032],
                ]
            }
        )
        start, end, _sample = fd._parse_info_sheet(workbook, 2026)
        assert start == date(2026, 1, 16)
        assert end == date(2026, 1, 19)

    def test_sample_size_a_fallback_numeric_in_col_b(self) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    ["Sample size", 2032],
                ]
            }
        )
        _start, _end, sample = fd._parse_info_sheet(workbook, 2026)
        assert sample == 2032

    def test_sample_size_a_fallback_skips_blank_column_first(self) -> None:
        # Column B is blank (None), so the scan must "continue" past it to
        # column C, which holds the real value.
        workbook = build_workbook(
            {
                "Info": [
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    ["Sample size", None, 2032],
                ]
            }
        )
        _start, _end, sample = fd._parse_info_sheet(workbook, 2026)
        assert sample == 2032

    def test_sample_size_a_fallback_digits_in_string(self) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    ["Sample size", "n=2,032"],
                ]
            }
        )
        _start, _end, sample = fd._parse_info_sheet(workbook, 2026)
        assert sample == 2032

    def test_sample_size_falls_back_to_tables_sheet_unweighted_row(self) -> None:
        workbook = build_workbook(
            {
                "Info": [[None, "Dates conducted", "16-19 Jan 2026"]],
                "Tables": [["Unweighted sample", 2032]],
            }
        )
        _start, _end, sample = fd._parse_info_sheet(workbook, 2026)
        assert sample == 2032

    def test_missing_fieldwork_raises(self) -> None:
        workbook = build_workbook({"Info": [[None, "Sample size", 2032]]})
        with pytest.raises(ValueError, match="Dates conducted not found in workbook"):
            fd._parse_info_sheet(workbook, 2026)

    def test_missing_fieldwork_uses_fieldwork_label_override(self) -> None:
        workbook = build_workbook({"Info": [[None, "Sample size", 2032]]})
        start, end, sample = fd._parse_info_sheet(
            workbook, 2026, fieldwork_label="20-22 Feb 2026"
        )
        assert start == date(2026, 2, 20)
        assert end == date(2026, 2, 22)
        assert sample == 2032

    def test_missing_sample_size_raises(self) -> None:
        workbook = build_workbook(
            {
                "Info": [[None, "Dates conducted", "16-19 Jan 2026"]],
                "Tables": [["Nothing here"]],
            }
        )
        with pytest.raises(ValueError, match="Sample size not found in workbook"):
            fd._parse_info_sheet(workbook, 2026)

    def test_missing_sample_size_and_no_tables_sheet_propagates(self) -> None:
        workbook = build_workbook(
            {"Info": [[None, "Dates conducted", "16-19 Jan 2026"]]}
        )
        with pytest.raises(ValueError, match="Could not find tables sheet"):
            fd._parse_info_sheet(workbook, 2026)

    def test_sample_size_bc_value_without_digits_falls_through_to_a_fallback(
        self,
    ) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    [None, "Sample size", "undisclosed"],
                    ["Sample size", 2032],
                ]
            }
        )
        _start, _end, sample = fd._parse_info_sheet(workbook, 2026)
        assert sample == 2032

    def test_sample_size_a_fallback_col_without_digits_falls_through_to_next_col(
        self,
    ) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    ["Sample size", "TBD", "n=2032"],
                ]
            }
        )
        _start, _end, sample = fd._parse_info_sheet(workbook, 2026)
        assert sample == 2032

    def test_sample_size_a_fallback_row_with_no_value_tries_next_matching_row(
        self,
    ) -> None:
        workbook = build_workbook(
            {
                "Info": [
                    [None, "Dates conducted", "16-19 Jan 2026"],
                    ["Sample size", None, None, None, None],
                    ["Sample size", 2032],
                ]
            }
        )
        _start, _end, sample = fd._parse_info_sheet(workbook, 2026)
        assert sample == 2032

    def test_sample_size_tables_fallback_row_without_a_number_tries_next_row(
        self,
    ) -> None:
        workbook = build_workbook(
            {
                "Info": [[None, "Dates conducted", "16-19 Jan 2026"]],
                "Tables": [
                    ["Unweighted sample", "n/a", "n/a", "n/a", "n/a", "n/a", "n/a"],
                    ["Unweighted sample", 2032],
                ],
            }
        )
        _start, _end, sample = fd._parse_info_sheet(workbook, 2026)
        assert sample == 2032


# ── _find_tables_sheet ───────────────────────────────────────────────────────


class TestFindTablesSheet:
    """Tests for _find_tables_sheet — locate the cross-tabulation sheet."""

    def test_matches_exact_name_tables_case_insensitively(self) -> None:
        workbook = build_workbook({"Cover": [[1]], "TABLES": [[2]]})
        assert fd._find_tables_sheet(workbook).title == "TABLES"

    def test_matches_name_containing_table(self) -> None:
        workbook = build_workbook({"Cover": [[1]], "VI Tables Q1": [[2]]})
        assert fd._find_tables_sheet(workbook).title == "VI Tables Q1"

    def test_matches_name_containing_voting_intention(self) -> None:
        workbook = build_workbook({"Cover": [[1]], "Voting Intention Data": [[2]]})
        assert fd._find_tables_sheet(workbook).title == "Voting Intention Data"

    def test_matches_name_containing_respondents(self) -> None:
        workbook = build_workbook({"Cover": [[1]], "All Respondents": [[2]]})
        assert fd._find_tables_sheet(workbook).title == "All Respondents"

    def test_missing_raises(self) -> None:
        workbook = build_workbook({"Cover": [[1]], "Data": [[2]]})
        with pytest.raises(ValueError, match="Could not find tables sheet"):
            fd._find_tables_sheet(workbook)


# ── _find_vi_section_start ───────────────────────────────────────────────────


class TestFindViSectionStart:
    """Tests for _find_vi_section_start — locate the VI question header row."""

    def test_phase_one_combined_voting_intention_excluding(self) -> None:
        rows = [
            _row({1: "cover"}),
            _row({1: "Combined voting intention (excluding don't knows)"}),
        ]
        workbook = build_workbook({"Tables": rows})
        assert fd._find_vi_section_start(workbook["Tables"]) == 2

    def test_phase_one_takes_priority_over_an_earlier_phase_three_row(self) -> None:
        # Row 2 alone would satisfy phase 3 ("general election" + "vote
        # for"); row 10 satisfies phase 1. Since phase 1 runs its own full
        # scan before phase 2/3 are tried at all, row 10 must win even
        # though row 2 appears earlier in the sheet.
        rows = [_row({}) for _ in range(9)]
        rows[1] = _row({1: "Would vote for in a general election tomorrow"})
        rows.append(_row({1: "Combined voting intention (excluding don't knows)"}))
        workbook = build_workbook({"Tables": rows})
        assert fd._find_vi_section_start(workbook["Tables"]) == 10

    def test_phase_two_general_election_next_few_weeks(self) -> None:
        rows = [
            _row({1: "cover"}),
            _row({1: "If there were a general election held in the next few weeks"}),
        ]
        workbook = build_workbook({"Tables": rows})
        assert fd._find_vi_section_start(workbook["Tables"]) == 2

    def test_phase_three_general_election_and_vote_for(self) -> None:
        rows = [
            _row({1: "cover"}),
            _row({1: "Which party would you vote for in a general election"}),
        ]
        workbook = build_workbook({"Tables": rows})
        assert fd._find_vi_section_start(workbook["Tables"]) == 2

    def test_missing_raises(self) -> None:
        rows = [_row({1: "nothing relevant here"})]
        workbook = build_workbook({"Tables": rows})
        with pytest.raises(
            ValueError, match="Could not find voting intention section"
        ):
            fd._find_vi_section_start(workbook["Tables"])


# ── _parse_party_macro_percentages ─────────────────────────────────────────────


class TestParsePartyMacroPercentages:
    """Tests for _parse_party_macro_percentages — national and macro-region VI."""

    def test_happy_path_national_and_macro_values(self) -> None:
        workbook = build_workbook({"Tables": _tables_sheet()})
        parsed = fd._parse_party_macro_percentages(workbook)

        assert len(parsed) == 8
        assert parsed["Labour"][fd.NATIONAL_KEY] == 32.0
        assert parsed["Labour"]["Greater London"] == 45.0
        assert parsed["Labour"]["North of England"] == 38.0
        assert parsed["Labour"]["Scotland"] == 5.0
        assert parsed["Conservative"]["Greater London"] == 18.0
        assert parsed["Conservative"]["Scotland"] == 15.0
        # "Other" is the last party appended by _tables_sheet, so it is the
        # one most exposed to the sheet's-actual-last-row bug (see
        # test_party_on_sheets_actual_last_row_is_silently_dropped_pins_current_behaviour
        # below): without the mandatory trailer row, this would silently
        # read 0.0 (the optional-party default) instead of the real 2.0.
        assert parsed["Other"][fd.NATIONAL_KEY] == 2.0

    def test_party_on_sheets_actual_last_row_is_silently_dropped_pins_current_behaviour(
        self,
    ) -> None:
        """Pin a latent bug in ``_parse_party_macro_percentages``.

        ``focaldata_import.py:624`` scans
        ``range(start_row, min(sheet.max_row, start_row + 180))`` — an
        EXCLUSIVE upper bound — so whatever row sits at the sheet's actual
        ``max_row`` is never read at all. Here "Green" (a required party,
        not one of the three parties the code defaults to zero) is placed
        last with no trailer row after it, so its otherwise-valid row is
        silently skipped and the parse raises "Missing expected party rows"
        even though every party's data is genuinely present on the sheet —
        it is simply on the excluded last row. Latent: every real Focaldata
        sheet (polls 208-220 in the live DB) has a trailer row after the
        data, so this has not been observed in production. Not fixed here,
        per this piece's scope.
        """
        ordered = dict(_NATIONAL_FRACTIONS)
        green_fraction = ordered.pop("Green")
        ordered["Green"] = green_fraction  # re-insert last, with no trailer after
        rows = _header_rows()
        for party, fraction in ordered.items():
            rows.append(_party_row(party, fraction))
        # Deliberately NO trailer row here — that omission is the point.
        workbook = build_workbook({"Tables": rows})
        with pytest.raises(ValueError, match="Missing expected party rows") as exc_info:
            fd._parse_party_macro_percentages(workbook)
        assert "Green" in str(exc_info.value)

    def test_party_without_macro_cells_defaults_every_macro_to_zero(self) -> None:
        workbook = build_workbook({"Tables": _tables_sheet()})
        parsed = fd._parse_party_macro_percentages(workbook)
        # Green has a national row but no macro figures in _PARTY_MACROS.
        assert parsed["Green"]["North of England"] == 0.0
        assert parsed["Green"]["Scotland"] == 0.0

    def test_region_header_not_found_leaves_only_national_key(self) -> None:
        workbook = build_workbook({"Tables": _tables_sheet(macro_columns=None)})
        parsed = fd._parse_party_macro_percentages(workbook)
        assert parsed["Labour"] == {fd.NATIONAL_KEY: 32.0}

    def test_unmapped_region_header_column_is_ignored(self) -> None:
        header = {17: "Atlantis", **{col: name for name, col in _MACRO_COLUMNS.items()}}
        rows = _header_rows(macro_columns=None)
        rows.append(_row(header))
        for party, fraction in _NATIONAL_FRACTIONS.items():
            rows.append(_party_row(party, fraction, _PARTY_MACROS.get(party)))
        rows.append(_row({1: _TRAILER_ROW_TEXT}))

        workbook = build_workbook({"Tables": rows})
        parsed = fd._parse_party_macro_percentages(workbook)
        assert set(parsed["Labour"].keys()) == {fd.NATIONAL_KEY, *_MACRO_COLUMNS}
        assert parsed["Labour"]["Greater London"] == 45.0

    def test_canonical_from_col1_takes_priority_over_col2(self) -> None:
        # Col 1 is a canonical party ("Labour"); col 2 also happens to be a
        # canonical label ("Green") but must be ignored entirely once col 1
        # has already matched, and the value must fall back to col 3.
        #
        # Green's own genuine row comes *before* the ambiguous row, so if the
        # "canonical is None" priority guard were ever dropped, the col2
        # branch would fire on the ambiguous row and overwrite Green's
        # already-parsed 6.0 with the wrong value (32, from col 3) — a
        # witness inside the mutation's reach, not one a later legitimate
        # row would quietly restore.
        rows = _header_rows()
        rows.append(_party_row("Green", _NATIONAL_FRACTIONS["Green"]))
        rows.append(_row({1: "Labour", 2: "Green", 3: 0.32}))
        for party in _REQUIRED_PARTIES:
            if party in {"Labour", "Green"}:
                continue
            rows.append(_party_row(party, _NATIONAL_FRACTIONS[party]))
        rows.append(_row({1: _TRAILER_ROW_TEXT}))
        workbook = build_workbook({"Tables": rows})
        parsed = fd._parse_party_macro_percentages(workbook)
        assert parsed["Labour"][fd.NATIONAL_KEY] == 32.0
        assert parsed["Green"][fd.NATIONAL_KEY] == 6.0

    def test_col1_value_prefers_col2_over_col3_when_col2_numeric(self) -> None:
        rows = _header_rows()
        rows.append(_row({1: "Labour", 2: 0.32, 3: 0.99}))
        for party in _REQUIRED_PARTIES:
            if party == "Labour":
                continue
            rows.append(_party_row(party, _NATIONAL_FRACTIONS[party]))
        rows.append(_row({1: _TRAILER_ROW_TEXT}))
        workbook = build_workbook({"Tables": rows})
        parsed = fd._parse_party_macro_percentages(workbook)
        assert parsed["Labour"][fd.NATIONAL_KEY] == 32.0

    def test_col1_value_falls_back_to_col3_when_col2_not_numeric(self) -> None:
        rows = _header_rows()
        rows.append(_row({1: "Labour", 2: None, 3: 0.32}))
        for party in _REQUIRED_PARTIES:
            if party == "Labour":
                continue
            rows.append(_party_row(party, _NATIONAL_FRACTIONS[party]))
        rows.append(_row({1: _TRAILER_ROW_TEXT}))
        workbook = build_workbook({"Tables": rows})
        parsed = fd._parse_party_macro_percentages(workbook)
        assert parsed["Labour"][fd.NATIONAL_KEY] == 32.0

    def test_canonical_from_col2_used_when_col1_unmatched(self) -> None:
        rows = _header_rows()
        rows.append(_row({1: "Party:", 2: "Labour", 3: 0.32}))
        for party in _REQUIRED_PARTIES:
            if party == "Labour":
                continue
            rows.append(_party_row(party, _NATIONAL_FRACTIONS[party]))
        rows.append(_row({1: _TRAILER_ROW_TEXT}))
        workbook = build_workbook({"Tables": rows})
        parsed = fd._parse_party_macro_percentages(workbook)
        assert parsed["Labour"][fd.NATIONAL_KEY] == 32.0

    def test_canonical_from_col2_value_stays_none_when_neither_cell_numeric(
        self,
    ) -> None:
        # Exercises the col2-branch's own value fallback
        # (focaldata_import.py:656-658): col 3 isn't numeric, so the code
        # falls back to re-reading col 2's own cell — but col 2 is occupied
        # by the party-name label text itself ("Green"), never numeric, so
        # `value` stays None and this ambiguous row contributes nothing.
        # Green's real value must come from its own genuine row just after.
        rows = _header_rows()
        rows.append(_row({1: "Party:", 2: "Green", 3: "n/a"}))
        rows.append(_party_row("Green", _NATIONAL_FRACTIONS["Green"]))
        for party in _REQUIRED_PARTIES:
            if party == "Green":
                continue
            rows.append(_party_row(party, _NATIONAL_FRACTIONS[party]))
        rows.append(_row({1: _TRAILER_ROW_TEXT}))
        workbook = build_workbook({"Tables": rows})
        parsed = fd._parse_party_macro_percentages(workbook)
        assert parsed["Green"][fd.NATIONAL_KEY] == 6.0

    def test_row_not_matching_any_party_is_skipped(self) -> None:
        rows = _header_rows()
        rows.append(_row({1: "Notes: totals may not sum to 100%", 2: 42.0}))
        for party in _REQUIRED_PARTIES:
            rows.append(_party_row(party, _NATIONAL_FRACTIONS[party]))
        rows.append(_row({1: _TRAILER_ROW_TEXT}))
        workbook = build_workbook({"Tables": rows})
        parsed = fd._parse_party_macro_percentages(workbook)
        assert len(parsed) == 8
        # Confirms the trailer row actually made "Other" (the last party
        # appended above) reach real parsing, not the optional-party 0.0
        # default — otherwise this len()==8 check alone would prove nothing.
        assert parsed["Other"][fd.NATIONAL_KEY] == 2.0

    def test_neither_reading_valid_skips_party_and_raises(self) -> None:
        national = dict(_NATIONAL_FRACTIONS)
        del national["Labour"]
        rows = _header_rows()
        rows.append(_row({1: "Labour", 2: None, 3: None}))
        for party, fraction in national.items():
            rows.append(_party_row(party, fraction))
        rows.append(_row({1: _TRAILER_ROW_TEXT}))
        workbook = build_workbook({"Tables": rows})
        with pytest.raises(ValueError, match="Missing expected party rows") as exc_info:
            fd._parse_party_macro_percentages(workbook)
        assert "Labour" in str(exc_info.value)

    def test_column_n_prefix_stops_national_scan(self) -> None:
        national_before = {"Conservative": 0.24, "Labour": 0.32}
        national_after = {
            name: pct
            for name, pct in _NATIONAL_FRACTIONS.items()
            if name not in national_before
        }
        extra = [_row({1: "Column N=2032"})]
        for party, fraction in national_after.items():
            extra.append(_party_row(party, fraction))
        workbook = build_workbook(
            {"Tables": _tables_sheet(national=national_before, extra_rows=extra)}
        )
        with pytest.raises(ValueError, match="Missing expected party rows") as exc_info:
            fd._parse_party_macro_percentages(workbook)
        message = str(exc_info.value)
        assert "Conservative" not in message
        assert "Labour" not in message
        assert "Green" in message

    def test_q_prefix_at_row_start_plus_three_does_not_break(self) -> None:
        # The production check is `row > start_row + 3`. Placing the "Q1.
        # ..." row at exactly start_row + 3 (row 4, since the header phrase
        # is row 1 = start_row) means it must NOT break: `4 > 4` is False.
        # A mutant that changes the "+3" to "+2" would make `4 > 3` True and
        # incorrectly break here — pinning the row at the exact boundary
        # (rather than anywhere "early") is what catches that mutant.
        rows = _header_rows()
        rows.append(_party_row("Conservative", _NATIONAL_FRACTIONS["Conservative"]))
        rows.append(_row({1: "Q1. Which party would you vote for?"}))
        for party, fraction in _NATIONAL_FRACTIONS.items():
            if party == "Conservative":
                continue
            rows.append(_party_row(party, fraction))
        rows.append(_row({1: _TRAILER_ROW_TEXT}))
        workbook = build_workbook({"Tables": rows})
        parsed = fd._parse_party_macro_percentages(workbook)
        assert len(parsed) == 8

    def test_q_prefix_at_row_start_plus_four_stops_scan(self) -> None:
        # The smallest row where `row > start_row + 3` is True: row 5 (since
        # start_row is row 1). Placed here, the "Q2. ..." row must break the
        # scan, so the parties after it are missing.
        national_before = {"Conservative": 0.24, "Labour": 0.32}
        national_after = {
            name: pct
            for name, pct in _NATIONAL_FRACTIONS.items()
            if name not in national_before
        }
        rows = _header_rows()
        for party, fraction in national_before.items():
            rows.append(_party_row(party, fraction))
        rows.append(_row({1: "Q2. Second question here"}))
        for party, fraction in national_after.items():
            rows.append(_party_row(party, fraction))
        rows.append(_row({1: _TRAILER_ROW_TEXT}))
        workbook = build_workbook({"Tables": rows})
        with pytest.raises(ValueError, match="Missing expected party rows") as exc_info:
            fd._parse_party_macro_percentages(workbook)
        message = str(exc_info.value)
        assert "Conservative" not in message
        assert "Labour" not in message
        assert "Green" in message

    def test_missing_non_optional_party_raises(self) -> None:
        national = {
            name: pct for name, pct in _NATIONAL_FRACTIONS.items() if name != "Green"
        }
        workbook = build_workbook({"Tables": _tables_sheet(national=national)})
        with pytest.raises(ValueError, match="Missing expected party rows") as exc_info:
            fd._parse_party_macro_percentages(workbook)
        assert "Green" in str(exc_info.value)

    def test_missing_multiple_non_optional_parties_lists_both_sorted(self) -> None:
        national = {
            name: pct
            for name, pct in _NATIONAL_FRACTIONS.items()
            if name not in {"Conservative", "Green"}
        }
        workbook = build_workbook({"Tables": _tables_sheet(national=national)})
        with pytest.raises(ValueError, match="Missing expected party rows") as exc_info:
            fd._parse_party_macro_percentages(workbook)
        message = str(exc_info.value)
        assert message.index("Conservative") < message.index("Green")

    @pytest.mark.parametrize(
        "party_name", ["Scottish National Party", "Plaid Cymru", "Other"]
    )
    def test_missing_optional_party_defaults_to_zero(self, party_name: str) -> None:
        # Contrast with bmg_research_import.py's
        # _parse_party_region_percentages: there, the required-party check
        # compares against the raw parsed rows rather than the defaulted
        # dict, so a workbook missing one of the three "optional" parties
        # wrongly raises. Here, the check runs against `parsed` *after* the
        # optional-party defaulting loop has already run `setdefault` on it,
        # so a workbook missing one of these three genuinely does NOT raise
        # — this matches the docstring and is not a bug.
        national = {
            name: pct for name, pct in _NATIONAL_FRACTIONS.items() if name != party_name
        }
        workbook = build_workbook({"Tables": _tables_sheet(national=national)})
        parsed = fd._parse_party_macro_percentages(workbook)
        assert parsed[party_name][fd.NATIONAL_KEY] == 0.0
        for macro_name in fd.MACRO_TO_INTERNAL_REGIONS:
            assert parsed[party_name][macro_name] == 0.0


# ── parse_poll ────────────────────────────────────────────────────────────────


class TestParsePoll:
    """Tests for parse_poll — combines fieldwork, sample and VI parsing."""

    def test_happy_path(self) -> None:
        workbook = _full_workbook()
        parsed = fd.parse_poll(workbook, source_url=_XLSX_URL)

        assert parsed.sample_size == 2032
        assert parsed.fieldwork_start == date(2026, 1, 16)
        assert parsed.fieldwork_end == date(2026, 1, 19)
        assert len(parsed.party_macro_percentages) == 8
        assert parsed.party_macro_percentages["Labour"][fd.NATIONAL_KEY] == 32.0

    def test_year_hint_used_when_url_and_text_have_no_year(self) -> None:
        workbook = _full_workbook(fieldwork_text="16-19 Jan")
        parsed = fd.parse_poll(
            workbook,
            source_url="https://focaldata.example.test/tables.xlsx",
            year_hint=2026,
        )
        assert parsed.fieldwork_start == date(2026, 1, 16)

    def test_fieldwork_label_used_when_info_sheet_has_none(self) -> None:
        workbook = build_workbook(
            {
                "Info": _info_sheet(fieldwork_label=None),
                "Tables": _tables_sheet(),
            }
        )
        parsed = fd.parse_poll(
            workbook, source_url=_XLSX_URL, fieldwork_label="20-22 Jan 2026"
        )
        assert parsed.fieldwork_start == date(2026, 1, 20)
        assert parsed.fieldwork_end == date(2026, 1, 22)

    def test_no_year_and_no_hint_raises(self) -> None:
        workbook = _full_workbook(fieldwork_text="16-19 Jan")
        with pytest.raises(ValueError, match="Could not infer year"):
            fd.parse_poll(
                workbook, source_url="https://focaldata.example.test/tables.xlsx"
            )

    def test_propagates_missing_party_error(self) -> None:
        national = {
            name: pct for name, pct in _NATIONAL_FRACTIONS.items() if name != "Green"
        }
        workbook = _full_workbook(national=national)
        with pytest.raises(ValueError, match="Missing expected party rows"):
            fd.parse_poll(workbook, source_url=_XLSX_URL)


# ── build_import_plan ──────────────────────────────────────────────────────────


class TestBuildImportPlan:
    """Tests for build_import_plan — dry-run plan from a fetched workbook."""

    def test_map_not_found_raises(
        self, db: Database, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        # A map-not-found plan must never reach the network: prove it by
        # making any fetch attempt fail the test outright.
        def fail_extract_workbook(*_a: object, **_k: object) -> Workbook:
            raise AssertionError("must not fetch when the map lookup already failed")

        monkeypatch.setattr(fd, "extract_workbook", fail_extract_workbook)

        with pytest.raises(ValueError, match="Map not found: 'No Such Map'"):
            fd.build_import_plan(db, map_name="No Such Map")

    def test_missing_parties_raises(
        self, db: Database, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        poll_map = db.add_map("A Map", parliament="westminster")

        def fake_extract_workbook(*_a: object, **_k: object) -> Workbook:
            return _full_workbook()

        monkeypatch.setattr(fd, "extract_workbook", fake_extract_workbook)

        with pytest.raises(ValueError, match="Missing parties in database"):
            fd.build_import_plan(db, map_name=poll_map.name)

    def test_macro_to_region_expansion(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        world = westminster_world

        def fake_extract_workbook(*_a: object, **_k: object) -> Workbook:
            return _full_workbook()

        monkeypatch.setattr(fd, "extract_workbook", fake_extract_workbook)

        plan = fd.build_import_plan(
            db,
            xlsx_url=_XLSX_URL,
            map_name=world.map_name,
            pollster_identifier="focaldata",
        )

        assert plan.map_id == world.map_id
        assert plan.map_name == world.map_name
        assert plan.pollster_exists is False
        assert plan.pollster_id is None
        assert plan.pollster_name == "Focaldata"
        assert plan.regions_mapping == ""
        assert plan.poll_exists is False
        assert plan.poll_id is None

        # 8 parties × (1 national row + 12 regional rows).
        assert len(plan.rows) == 8 * 13
        national_rows = [row for row in plan.rows if row.region_id is None]
        assert len(national_rows) == 8

        def _row_for(party: str, region: str | None) -> fd.PlannedPollRow:
            if region is None:
                return next(
                    r for r in plan.rows if r.party_name == party and r.region_id is None
                )
            return next(
                r
                for r in plan.rows
                if r.party_name == party and r.region_id == world.region_ids[region]
            )

        labour_national = _row_for("Labour", None)
        assert labour_national.percentage == 32.0
        assert labour_national.region_name == "National"
        assert labour_national.party_id == world.party_ids["Labour"]

        # "Greater London" maps to exactly one internal region: 1:1 applied.
        assert _row_for("Labour", "London").percentage == 45.0

        # "North of England" maps to three internal regions: the same macro
        # figure is replicated across all three.
        assert _row_for("Labour", "North East England").percentage == 38.0
        assert _row_for("Labour", "North West England").percentage == 38.0
        assert _row_for("Labour", "Yorkshire and The Humber").percentage == 38.0

        # A second, distinct party's own figures, so a mutant reusing one
        # party's percentages for every row fails.
        assert _row_for("Conservative", "London").percentage == 18.0
        assert _row_for("Conservative", "Scotland").percentage == 15.0

    def test_region_outside_every_macro_gets_zero_row_pins_current_behaviour(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """Pin the documented, already-known "region outside every macro" behaviour.

        "Northern Ireland" is not the target of any macro in
        MACRO_TO_INTERNAL_REGIONS, so every DB region named "Northern
        Ireland" gets an explicit 0.0 row for every party, however large
        that party's other regional or national figures are. This is a DB
        region that sits outside every macro entirely — a different case
        from a *mapped* internal region that's simply missing from a given
        DB map (see test_internal_region_absent_from_map_does_not_raise
        below, where Opinium/Ipsos/etc. raise instead). Not fixed in this
        piece.
        """
        world = westminster_world
        # Every party (including the three "optional" ones) gets the same
        # non-zero figure in every one of the six macros, so if any macro's
        # value leaked into Northern Ireland's row it would show up here as
        # non-zero.
        macros_all_parties = {
            party: {macro: 0.5 for macro in fd.MACRO_TO_INTERNAL_REGIONS}
            for party in _NATIONAL_FRACTIONS
        }

        def fake_extract_workbook(*_a: object, **_k: object) -> Workbook:
            return build_workbook(
                {
                    "Info": _info_sheet(),
                    "Tables": _tables_sheet(macros=macros_all_parties),
                }
            )

        monkeypatch.setattr(fd, "extract_workbook", fake_extract_workbook)

        plan = fd.build_import_plan(
            db,
            xlsx_url=_XLSX_URL,
            map_name=world.map_name,
            pollster_identifier="focaldata",
        )
        ni_rows = [
            row for row in plan.rows if row.region_id == world.region_ids["Northern Ireland"]
        ]
        assert len(ni_rows) == 8
        assert all(row.percentage == 0.0 for row in ni_rows)

    def test_internal_region_absent_from_map_does_not_raise(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        # The mirror-image case of the Northern Ireland test above: here the
        # DB's map is simply missing most internal regions altogether (only
        # "London" and "Scotland" are seeded on this second map), so most of
        # MACRO_TO_INTERNAL_REGIONS' internal names — regions the macros DO
        # map to — never resolve to a region id at all. Unlike
        # Opinium/Ipsos/Lord Ashcroft/YouGov/Deltapoll, which raise when a
        # *mapped* internal region is missing from the DB map, Focaldata
        # must not raise or fabricate a row here — it just leaves those
        # macros' region-id lists empty. (This does not exercise the
        # `region_id is not None` guard specifically — removing that guard
        # would append `None` to the id list, which never matches a real
        # region.id either, so this test cannot tell the guard apart from
        # its absence; it only pins the observable "no raise, no phantom
        # row" outcome.)
        small_map = db.add_map("Focaldata Test Map", parliament="westminster")
        london_id = db.add_region(small_map.id, "London").id
        scotland_id = db.add_region(small_map.id, "Scotland").id

        def fake_extract_workbook(*_a: object, **_k: object) -> Workbook:
            return _full_workbook()

        monkeypatch.setattr(fd, "extract_workbook", fake_extract_workbook)

        plan = fd.build_import_plan(
            db,
            xlsx_url=_XLSX_URL,
            map_name=small_map.name,
            pollster_identifier="focaldata",
        )

        region_row_ids = {row.region_id for row in plan.rows if row.region_id is not None}
        assert region_row_ids == {london_id, scotland_id}
        # 8 parties × (1 national + 2 regional) rows.
        assert len(plan.rows) == 8 * 3

        labour_london = next(
            r for r in plan.rows if r.party_name == "Labour" and r.region_id == london_id
        )
        assert labour_london.percentage == 45.0

    def test_pollster_and_poll_existing_flags(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        world = westminster_world
        # Named differently from the code's "Focaldata" fallback default, so
        # a mutant that always returns the hard-coded default cannot pass.
        pollster = db.add_pollster("Focaldata Research", "focaldata", weight=1.0)
        poll = db.add_poll(
            pollster.id,
            world.map_id,
            date(2026, 1, 16),
            date(2026, 1, 19),
            sample_size=2032,
        )

        def fake_extract_workbook(*_a: object, **_k: object) -> Workbook:
            return _full_workbook()

        monkeypatch.setattr(fd, "extract_workbook", fake_extract_workbook)

        plan = fd.build_import_plan(
            db,
            xlsx_url=_XLSX_URL,
            map_name=world.map_name,
            pollster_identifier="focaldata",
        )

        assert plan.pollster_exists is True
        assert plan.pollster_id == pollster.id
        assert plan.pollster_name == "Focaldata Research"
        assert plan.poll_exists is True
        assert plan.poll_id == poll.id


# ── _cli_preview ────────────────────────────────────────────────────────────


class TestCliPreview:
    """Tests for _cli_preview — dry-run summary printed to stdout."""

    def _plan(self, **overrides: Any) -> fd.ImportPlan:
        rows = overrides.pop(
            "rows",
            [
                fd.PlannedPollRow(
                    party_id=1,
                    party_name="Labour",
                    region_id=None,
                    region_name="National",
                    percentage=32.0,
                ),
                fd.PlannedPollRow(
                    party_id=1,
                    party_name="Labour",
                    region_id=3,
                    region_name="London",
                    percentage=45.0,
                ),
            ],
        )
        defaults: dict[str, Any] = {
            "pollster_identifier": "focaldata",
            "pollster_name": "Focaldata",
            "pollster_id": None,
            "pollster_exists": False,
            "regions_mapping": "",
            "map_id": 1,
            "map_name": "UK Constituencies post 2022",
            "source_url": _XLSX_URL,
            "parsed": fd.ParsedPoll(
                sample_size=2032,
                fieldwork_start=date(2026, 1, 16),
                fieldwork_end=date(2026, 1, 19),
                party_macro_percentages={},
            ),
            "poll_id": None,
            "poll_exists": False,
            "rows": rows,
        }
        defaults.update(overrides)
        return fd.ImportPlan(**defaults)

    def test_new_pollster_and_poll(self, capsys: pytest.CaptureFixture[str]) -> None:
        fd._cli_preview(self._plan())
        out = capsys.readouterr().out

        assert "Parsed poll: fieldwork=2026-01-16 to 2026-01-19, sample=2032" in out
        assert "[dry-run] would create pollster: focaldata" in out
        # "in out" would also match the pollster line above (a prefix of
        # it), so check the exact printed line instead.
        assert "[dry-run] would create poll" in out.splitlines()
        assert (
            "[dry-run] would insert row: party=Labour, region=National, "
            "region_id=None, pct=32.00" in out
        )
        assert (
            "[dry-run] would insert row: party=Labour, region=London, "
            "region_id=3, pct=45.00" in out
        )

    def test_existing_pollster_and_poll(
        self, capsys: pytest.CaptureFixture[str]
    ) -> None:
        fd._cli_preview(self._plan(pollster_exists=True, poll_exists=True, poll_id=9))
        out = capsys.readouterr().out

        assert "pollster exists: focaldata" in out
        assert "poll exists: 9" in out
        assert "[dry-run] would create pollster" not in out
        assert "[dry-run] would create poll" not in out


# ── main ────────────────────────────────────────────────────────────────────


class TestMain:
    """Tests for main — the CLI entry point, run against the ``db`` fixture."""

    def _run(
        self,
        db: Database,
        monkeypatch: pytest.MonkeyPatch,
        workbook: Workbook,
        *argv: str,
        fetched_urls: list[str] | None = None,
    ) -> None:
        """Run main() with Database and extract_workbook faked, via sys.argv.

        If ``fetched_urls`` is given, each URL ``extract_workbook`` is called
        with is appended to it, so a test can prove ``--xlsx-url`` actually
        reaches the fetch rather than the CLI default silently being used.
        """

        def fake_database(*_a: object, **_k: object) -> Database:
            return db

        def fake_extract_workbook(xlsx_url: str) -> Workbook:
            if fetched_urls is not None:
                fetched_urls.append(xlsx_url)
            return workbook

        monkeypatch.setattr(fd, "Database", fake_database)
        monkeypatch.setattr(fd, "extract_workbook", fake_extract_workbook)
        monkeypatch.setattr(sys, "argv", ["focaldata_import.py", *argv])
        fd.main()

    def test_dry_run_previews_without_writing(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        world = westminster_world
        self._run(
            db, monkeypatch, _full_workbook(), "--map-name", world.map_name, "--dry-run"
        )

        out = capsys.readouterr().out
        assert f"Fetching XLSX: {fd.DEFAULT_XLSX_URL}" in out
        assert "[dry-run] would create pollster: focaldata" in out
        assert len(db.get_all_pollsters()) == 0

    def test_commit_creates_pollster_poll_and_rows(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        world = westminster_world
        fetched_urls: list[str] = []
        self._run(
            db,
            monkeypatch,
            _full_workbook(),
            "--map-name",
            world.map_name,
            "--xlsx-url",
            "https://focaldata.example.test/2026/other.xlsx",
            fetched_urls=fetched_urls,
        )

        out = capsys.readouterr().out
        assert "created pollster: focaldata" in out
        assert "inserted poll rows: 104" in out

        pollsters = db.get_all_pollsters()
        assert len(pollsters) == 1
        assert pollsters[0].identifier == "focaldata"
        polls = db.get_polls_by_pollster(pollsters[0].id)
        assert len(polls) == 1
        assert len(db.get_rows_for_poll(polls[0].id)) == 104
        assert fetched_urls == ["https://focaldata.example.test/2026/other.xlsx"]
        assert polls[0].source_url == "https://focaldata.example.test/2026/other.xlsx"

    def test_rerun_skips_then_replace_rows_overwrites(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        world = westminster_world
        argv = ("--map-name", world.map_name)

        self._run(db, monkeypatch, _full_workbook(), *argv)
        pollster = db.get_pollster_by_identifier("focaldata")
        assert pollster is not None
        poll_id = db.get_polls_by_pollster(pollster.id)[0].id
        assert len(db.get_rows_for_poll(poll_id)) == 104
        capsys.readouterr()  # drain the first run's output before the second

        self._run(db, monkeypatch, _full_workbook(), *argv)
        out = capsys.readouterr().out
        assert f"poll exists: {poll_id}" in out.splitlines()
        assert f"poll {poll_id} already has rows; use --replace-rows" in out
        assert len(db.get_rows_for_poll(poll_id)) == 104

        self._run(db, monkeypatch, _full_workbook(), *argv, "--replace-rows")
        out = capsys.readouterr().out
        assert "deleted existing rows: 104" in out
        assert "inserted poll rows: 104" in out
        assert len(db.get_rows_for_poll(poll_id)) == 104

    def test_map_name_cli_flag_reaches_build_import_plan(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        # --map-name must reach build_import_plan rather than silently
        # falling back to DEFAULT_MAP_NAME (or westminster_world's map,
        # which happens to equal the default — see test_cli_defaults_pin):
        # seed a second, differently-named map (parties are not map-scoped,
        # so westminster_world's parties are still available) and assert
        # the committed poll's map_id is the second map's.
        world = westminster_world
        other_map = db.add_map("Focaldata Second Map", parliament="westminster")
        db.add_region(other_map.id, "London")

        self._run(db, monkeypatch, _full_workbook(), "--map-name", other_map.name)

        pollster = db.get_pollster_by_identifier("focaldata")
        assert pollster is not None
        poll = db.get_polls_by_pollster(pollster.id)[0]
        assert poll.map_id == other_map.id
        assert poll.map_id != world.map_id

    def test_pollster_identifier_cli_flag_reaches_build_import_plan(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        # --pollster-identifier must reach build_import_plan rather than
        # silently falling back to DEFAULT_POLLSTER_IDENTIFIER.
        world = westminster_world
        self._run(
            db,
            monkeypatch,
            _full_workbook(),
            "--map-name",
            world.map_name,
            "--pollster-identifier",
            "focaldata_alt",
        )

        assert db.get_pollster_by_identifier("focaldata_alt") is not None
        assert db.get_pollster_by_identifier("focaldata") is None

    def test_year_hint_cli_flag(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        world = westminster_world
        workbook = _full_workbook(fieldwork_text="16-19 Jan")

        self._run(
            db,
            monkeypatch,
            workbook,
            "--map-name",
            world.map_name,
            "--xlsx-url",
            "https://focaldata.example.test/tables.xlsx",
            "--year-hint",
            "2026",
            "--dry-run",
        )

        out = capsys.readouterr().out
        assert "Parsed poll: fieldwork=2026-01-16 to 2026-01-19" in out

    def test_cli_defaults_pin(
        self,
        db: Database,
        westminster_world: WestminsterWorld,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        world = westminster_world
        assert fd.DEFAULT_MAP_NAME == world.map_name
        assert fd.DEFAULT_POLLSTER_IDENTIFIER == "focaldata"

        self._run(db, monkeypatch, _full_workbook(), "--dry-run")

        out = capsys.readouterr().out
        assert "[dry-run] would create pollster: focaldata" in out
